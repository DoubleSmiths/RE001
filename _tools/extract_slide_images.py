import sys, os, io
from pptx import Presentation
from pptx.util import Emu

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

pptx_path = sys.argv[1]
slide_num = int(sys.argv[2])
out_dir = sys.argv[3]

os.makedirs(out_dir, exist_ok=True)

prs = Presentation(pptx_path)
slide = prs.slides[slide_num - 1]

count = 0
for idx, shape in enumerate(slide.shapes):
    if shape.shape_type == 13:  # PICTURE
        try:
            img = shape.image
            ext = img.ext
            fname = f"slide{slide_num}_shape{idx:03d}_{shape.name.replace(' ','_').replace(';','_')}.{ext}"
            fname = fname.replace('/', '_')
            with open(os.path.join(out_dir, fname), "wb") as f:
                f.write(img.blob)
            print(f"{fname} ({shape.width}x{shape.height} EMU)")
            count += 1
        except Exception as e:
            print(f"ERR idx={idx}: {e}")
print(f"Total: {count}")
