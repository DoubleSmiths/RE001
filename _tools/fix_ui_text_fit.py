"""영문이 UI 칸을 넘치는 문자열을 짧은 표기로 바꾼다 — 캡처로 넘침을 확인한 것만 적는다.

  사용: python _tools/fix_ui_text_fit.py          # 미리보기
        python _tools/fix_ui_text_fit.py --write  # LocalizationData.xlsx 저장 → Unity Tools/ParseAllLocal
"""
import sys

import openpyxl

XLSX = "Assets/@Resources/Data/Excel/LocalizationData.xlsx"

# TextKey → (예전 영문, 새 영문, 근거)
OVERRIDES = {
    "ForgeDurabilityText": ("DURABILITY", "DUR.",
                            "대장간 수리 칸 'DURABILITY 110/110' 이 옆 칸(강화)까지 넘침 (1080x2340 캡처, 2026-09-29)"),
}


def main(write):
    wb = openpyxl.load_workbook(XLSX)
    changed = 0
    for ws in wb.worksheets:
        header = [c.value for c in ws[1]]
        if "TextKey" not in header:
            continue
        k, e = header.index("TextKey"), next(i for i, h in enumerate(header) if h and str(h).lower().startswith("eng"))
        for row in ws.iter_rows(min_row=2):
            key = row[k].value
            if key in OVERRIDES:
                old, new, why = OVERRIDES[key]
                cur = row[e].value
                if cur == new:
                    continue
                if cur != old:
                    print(f"SKIP {ws.title}!{key}: 예상 '{old}' 인데 '{cur}'")
                    continue
                print(f"{ws.title}!{key}: '{cur}' -> '{new}'  ({why})")
                row[e].value = new
                changed += 1
    if write and changed:
        wb.save(XLSX)
        print(f"저장 {changed}건 — Unity Tools/ParseAllLocal 로 JSON 재생성")
    elif not write:
        print(f"미리보기 {changed}건 — 저장하려면 --write")


if __name__ == "__main__":
    main("--write" in sys.argv)
