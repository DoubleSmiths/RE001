"""
Quest_Unlock R2(Q1 ITEM_REWARD) TargetID 9001 -> 20001
JSON과 일치시킴.
"""
from pathlib import Path
import openpyxl

XLSX = Path(r"C:\stopresent\Unity\RE001\Assets\@Resources\Data\Excel\QuestData.xlsx")
wb = openpyxl.load_workbook(XLSX, data_only=False)

unlock = wb["Quest_Unlock"]
print("[Quest_Unlock]")
for r in range(2, unlock.max_row + 1):
    if (unlock.cell(r, 1).value == 1
            and unlock.cell(r, 2).value == "ITEM_REWARD"
            and unlock.cell(r, 3).value == 9001):
        unlock.cell(r, 3).value = 20001
        print(f"  Q1 R{r}: TargetID 9001 -> 20001")
        break
else:
    print("  변경 대상 없음 (이미 적용?)")

wb.save(XLSX)
print(f"\nSaved: {XLSX}")
