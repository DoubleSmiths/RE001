"""
Update Assets/@Resources/Data/Excel/QuestData.xlsx so the rows match the
already-edited JSON in Assets/@Resources/Data/JsonData/.

Changes:
  Sheet "Quest"
    Q3 (R4): CompleteTarget 1001 -> 10001,  FollowItemID 1001 -> 10001
    Q4 (R5): CompleteType MULTI -> BTN_CLICK
             CompleteTarget (empty) -> "BTN_GUILD_QUEST"
             FollowItemID 1001 -> 10001          # JSON에 이미 10001
    Q5 (R6): FollowItemID 1002 -> 10002          # JSON에 이미 10002
    Q6 (R7): CompleteTarget 1002 -> 10002, FollowItemID 1002 -> 10002

  Sheet "Quest_Condition"
    Quest 4 행 두 개(R4, R5) 삭제
"""
from pathlib import Path
import openpyxl

XLSX = Path(r"C:\stopresent\Unity\RE001\Assets\@Resources\Data\Excel\QuestData.xlsx")

wb = openpyxl.load_workbook(XLSX, data_only=False)

# ── Quest 시트 ─────────────────────────────────────
quest = wb["Quest"]

# 헤더(R1)에서 컬럼 인덱스 동적 조회 (1-based)
header = {quest.cell(1, c).value: c for c in range(1, quest.max_column + 1)}
COL_ID         = header["QuestID"]
COL_TYPE       = header["CompleteType"]
COL_TARGET     = header["CompleteTarget"]
COL_FOLLOW_ID  = header["FollowItemID"]


def set_quest_row(quest_id: int, **fields):
    for r in range(2, quest.max_row + 1):
        if quest.cell(r, COL_ID).value == quest_id:
            for col_name, value in fields.items():
                col_idx = header[col_name]
                quest.cell(r, col_idx).value = value
            print(f"  Quest {quest_id} (R{r}) updated: {fields}")
            return
    raise RuntimeError(f"Quest {quest_id} 행을 찾을 수 없음")


print("[Quest 시트]")
set_quest_row(3, CompleteTarget=10001, FollowItemID=10001)
set_quest_row(4, CompleteType="BTN_CLICK", CompleteTarget="BTN_GUILD_QUEST",
                  FollowItemID=10001)
set_quest_row(5, FollowItemID=10002)
set_quest_row(6, CompleteTarget=10002, FollowItemID=10002)

# ── Quest_Condition 시트 ───────────────────────────
cond = wb["Quest_Condition"]
# 1-based 인덱스. 뒤에서부터 지워야 행 번호 흔들림 없음.
to_delete = []
for r in range(2, cond.max_row + 1):
    if cond.cell(r, 1).value == 4:
        to_delete.append(r)

print("\n[Quest_Condition 시트]")
for r in reversed(to_delete):
    snapshot = [cond.cell(r, c).value for c in range(1, cond.max_column + 1)]
    cond.delete_rows(r, 1)
    print(f"  R{r} 삭제: {snapshot}")

wb.save(XLSX)
print(f"\nSaved: {XLSX}")
