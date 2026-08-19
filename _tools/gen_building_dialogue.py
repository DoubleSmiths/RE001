# -*- coding: utf-8 -*-
"""대장간·상점 NPC 잡담 대사를 넣는다.

## 기존 방식(그대로 따른다)
  대화 그룹 = DialogueData.xlsx의 **시트 하나**. 시트명이 곧 GroupKey.
  열은 TextKey / Kor / Eng. 잡담 시트는 `randomSingleLine: true`로 한 줄만 뽑아 쓴다
  (AlchemyNormalText / GuildNormalText가 그 방식이다).

## 화법
  대장장이   — 짧고 무뚝뚝한 반말. 손이 바쁜 사람의 말투.
  상점 주인  — 밝은 장사꾼 존댓말. 사고파는 쪽으로 유도한다.
  수도원 사제 — 느리고 정중한 존댓말. 축성(전직)을 권한다.

사용: python _tools/gen_building_dialogue.py   → 이후 Unity에서 Tools/ParseAllLocal
"""
import os

import openpyxl

BASE = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
DLG = os.path.join(BASE, "Assets", "@Resources", "Data", "Excel", "DialogueData.xlsx")

GROUPS = {
    "ForgeNormalText": [
        ("불 앞은 뜨거워.\n용건 있으면 빨리 말해.",
         "It's hot by the fire.\nSay your piece quickly."),
        ("재료만 가져와.\n모양은 내가 잡아 줄 테니까.",
         "Just bring me the materials.\nI'll shape them."),
        ("쇠는 정직해.\n두드린 만큼만 단단해지지.",
         "Iron is honest.\nIt hardens exactly as much as you strike it."),
        ("그 무기, 날이 상했는데.\n손봐 줄까?",
         "That blade's gone dull.\nWant me to work on it?"),
        ("강화는 운이야.\n실패해도 나한테 화내지 마.",
         "Enhancement is luck.\nDon't take a failure out on me."),
        ("좋은 장비 하나가\n목숨을 두 번 살려.",
         "One good piece of gear\nsaves your life twice."),
    ],
    "ShopNormalText": [
        ("어서 오세요!\n오늘은 뭘 찾으시나요?",
         "Welcome!\nWhat are you looking for today?"),
        ("재료는 넉넉히 사 두세요.\n필요할 때는 꼭 모자라거든요.",
         "Buy plenty of materials.\nYou always run short right when you need them."),
        ("안 쓰는 물건은 저한테 파세요.\n값은 후하게 쳐 드릴게요… 조금은요.",
         "Sell me what you don't use.\nI'll pay well... a little well."),
        ("강화 주문서요?\n안쪽 선반에 있습니다.",
         "Enhancement scrolls?\nThey're on the inner shelf."),
        ("골드는 쓰라고 있는 겁니다.\n쌓아두면 그냥 쇳조각이죠.",
         "Gold is meant to be spent.\nPiled up, it's just scrap metal."),
        ("또 오세요!\n다음엔 더 좋은 물건 들여놓을게요.",
         "Come again!\nI'll have better stock next time."),
    ],
    "MonasteryNormalText": [
        ("먼 길 오셨군요.\n잠시 숨을 고르고 가세요.",
         "You've come a long way.\nRest and catch your breath a while."),
        ("축성을 받으시겠습니까?\n길이 하나 더 열릴 것입니다.",
         "Would you receive a blessing?\nAnother path will open to you."),
        ("힘은 쥐는 것이 아니라\n맡겨지는 것입니다.",
         "Strength is not seized.\nIt is entrusted."),
        ("당신이 벤 것들도\n한때는 살아 있었습니다.",
         "The things you have cut down\nwere once alive, too."),
        ("두려움은 죄가 아닙니다.\n멈추지 않으면 됩니다.",
         "Fear is no sin.\nOnly stopping is."),
        ("여기서는 무기를 내려놓으셔도 됩니다.",
         "You may set down your weapon here."),
    ],
}


def write_group(wb, key, rows):
    if key in wb.sheetnames:
        del wb[key]
    ws = wb.create_sheet(key)
    ws.append(["TextKey", "Kor", "Eng"])
    for i, (kor, eng) in enumerate(rows, 1):
        ws.append([i, kor, eng])


def main():
    wb = openpyxl.load_workbook(DLG)
    for key, rows in GROUPS.items():
        write_group(wb, key, rows)
    wb.save(DLG)
    print(f"대화 그룹 {len(GROUPS)}개 기록: {', '.join(GROUPS)}")
    print("→ Unity에서 Tools/ParseAllLocal 필요")


if __name__ == "__main__":
    main()
