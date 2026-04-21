"""
QuestData.xlsx — Quest 5 (연금공방 해금) 흐름 정리.

Sheet "Quest" Q5(R6):
  CompleteType  ANIM_END                -> BTN_CLICK
  CompleteTarget ANIM_ALCHEMY_UNLOCK    -> GUILD_EXIT_QUEST_VFX
  AnimationID   ANIM_ALCHEMY_UNLOCK     -> (empty)
  AnimBlocking  TRUE                    -> (empty)

Sheet "Quest_Notify" Q5:
  BtnID  BTN_EXIT  ->  GUILD_EXIT_QUEST_VFX

Sheet "Quest_Unlock" Q5(UI_BUTTON):
  TargetID  BTN_ALCHEMY  ->  UNLOCK_ALCHEMY_BUTTON
"""
from pathlib import Path
import openpyxl

XLSX = Path(r"C:\stopresent\Unity\RE001\Assets\@Resources\Data\Excel\QuestData.xlsx")
wb = openpyxl.load_workbook(XLSX, data_only=False)

# ── Quest 시트 ─────────────────────────────────────
quest = wb["Quest"]
header = {quest.cell(1, c).value: c for c in range(1, quest.max_column + 1)}
COL_ID = header["QuestID"]


def update_quest(quest_id: int, **fields):
    for r in range(2, quest.max_row + 1):
        if quest.cell(r, COL_ID).value == quest_id:
            for k, v in fields.items():
                quest.cell(r, header[k]).value = v
            print(f"  Quest {quest_id} (R{r}) <- {fields}")
            return
    raise RuntimeError(f"Quest {quest_id} 없음")


print("[Quest 시트]")
update_quest(5,
             CompleteType="BTN_CLICK",
             CompleteTarget="GUILD_EXIT_QUEST_VFX",
             AnimationID=None,
             AnimBlocking=None)

# ── Quest_Notify 시트 ──────────────────────────────
notify = wb["Quest_Notify"]
print("\n[Quest_Notify 시트]")
for r in range(2, notify.max_row + 1):
    if notify.cell(r, 1).value == 5 and notify.cell(r, 3).value == "BTN_EXIT":
        notify.cell(r, 3).value = "GUILD_EXIT_QUEST_VFX"
        print(f"  Q5 R{r}: BTN_EXIT -> GUILD_EXIT_QUEST_VFX")
        break

# ── Quest_Unlock 시트 ──────────────────────────────
unlock = wb["Quest_Unlock"]
print("\n[Quest_Unlock 시트]")
for r in range(2, unlock.max_row + 1):
    if (unlock.cell(r, 1).value == 5
            and unlock.cell(r, 2).value == "UI_BUTTON"
            and unlock.cell(r, 3).value == "BTN_ALCHEMY"):
        unlock.cell(r, 3).value = "UNLOCK_ALCHEMY_BUTTON"
        print(f"  Q5 R{r}: BTN_ALCHEMY -> UNLOCK_ALCHEMY_BUTTON")
        break

wb.save(XLSX)
print(f"\nSaved: {XLSX}")
