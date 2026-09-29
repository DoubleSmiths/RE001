"""영문 용어·표기 통일 — LocalizationData / DialogueData 의 Eng 칸만 고친다. 키(TextKey)와 한국어는 건드리지 않는다.

  1. 'minion(s)' → 게임 고유명 'Minimon'. UI 라벨은 이미 MINIMON(도감·나의 미니몬 등)인데 능력 설명 약 500행은
     'minion' 이라 한 화면에 두 이름이 섞였다. 영문 제목도 'Minimon Quest' 다. 복수형은 고유명처럼 그대로 'Minimon'
     (스토어 문구와 같은 표기). 대문자 라벨은 'MINIMON', 그 외는 'Minimon'.
  2. 'Breezing Wind Forest' → 'Breezewood Forest'. 첫 지역(산들바람 숲)이 여정 필드 버튼·지역 배너에선
     BREEZEWOOD, 하위 지도 15곳·마물 서식지·길드 대사에선 Breezing Wind Forest 로 두 이름이었다.
     나머지 29개 지역은 한 이름이다. 배너 폭에 맞는 짧은 쪽으로 통일.
  3. 문자 그대로 박힌 '\\n'(백슬래시+n) → 실제 줄바꿈. 길드 대사 한 줄이 화면에 '\\n' 으로 나왔다.

  사용: python _tools/normalize_english_terms.py           # 미리보기(파일별 건수 + 예시)
        python _tools/normalize_english_terms.py --write   # 엑셀 저장 → Unity Tools/ParseAllLocal 로 JSON 재생성
"""
import re
import sys

import openpyxl

EXCEL = "Assets/@Resources/Data/Excel/"
FILES = ["LocalizationData.xlsx", "DialogueData.xlsx"]   # 대사 시트는 그룹마다 한 장, 열은 TextKey/Kor/Eng
MINION = re.compile(r"\b(minion)(s)?\b", re.IGNORECASE)
ESCAPED_NEWLINE = re.compile(r"\\+n")


def fix(text):
    text = MINION.sub(lambda m: "MINIMON" if m.group(1).isupper() else "Minimon", text)
    text = text.replace("Breezing Wind Forest", "Breezewood Forest")
    return ESCAPED_NEWLINE.sub("\n", text)


def run(name, write):
    wb = openpyxl.load_workbook(EXCEL + name)
    changed = []
    for ws in wb.worksheets:
        header = [c.value for c in ws[1]]
        if "Eng" not in header:
            continue
        eng = header.index("Eng")
        for row in ws.iter_rows(min_row=2):
            cell = row[eng]
            if not isinstance(cell.value, str):
                continue
            new = fix(cell.value)
            if new != cell.value:
                changed.append((ws.title, row[0].value, cell.value, new))
                cell.value = new
    for sheet, key, old, new in changed[:8]:
        print(f"  {sheet}/{key}: {old[:70]!r} -> {new[:70]!r}")
    print(f"{name}: {len(changed)}칸")
    if write and changed:
        wb.save(EXCEL + name)
    return len(changed)


if __name__ == "__main__":
    assert fix("A minion and MINIONS") == "A Minimon and MINIMON"
    assert fix("noisy\\\\naround Breezing Wind Forest F2") == "noisy\naround Breezewood Forest F2"
    write = "--write" in sys.argv
    total = sum(run(name, write) for name in FILES)
    if write and total:
        print("저장 — Unity Tools/ParseAllLocal 로 JSON 재생성")
