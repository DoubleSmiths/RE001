# -*- coding: utf-8 -*-
# JSON(현재 진실) → 로컬 .xlsx 재생성 — 구글시트 폐기(2026-07), 로컬 엑셀 원천 전환용.
# 대상: PlayerStatus/Creature/Item(3시트)/Localization/Dialogue/Travel/Buff
# (추가 7종 Skill~Combo는 AI/gen_data_sheets.py 사용, QuestData.xlsx는 이미 동기화됨)
# 컬럼 순서는 DataTransformer.cs의 각 로컬 파서와 1:1 — 파서를 바꾸면 여기도 같이 바꿀 것.
import json, os
from openpyxl import Workbook, load_workbook

ROOT = os.path.join(os.path.dirname(__file__), "..", "Assets", "@Resources", "Data")
JS, XL = os.path.join(ROOT, "JsonData"), os.path.join(ROOT, "Excel")

EQUIP_GRADE = ["Basic", "Abnormal", "Rare", "Exceptional", "Legendary", "Mystic", "REQUISITE"]

def j(name):
    with open(os.path.join(JS, name), encoding="utf-8-sig") as f:
        return json.load(f)

def cell(v):
    if v is None: return ""
    if isinstance(v, bool): return "TRUE" if v else "FALSE"
    if isinstance(v, list): return ",".join(str(x) for x in v)
    return v

def save(wb, name):
    wb.save(os.path.join(XL, name))
    print("생성:", name)

# ── PlayerStatusData.xlsx (3행 헤더, 파서 컬럼 순서: 기본13 + StartSkill 2 + 신규 스탯 10) ──
COLS_PS = ["DataId","Class","ClassProficiency","AdvRank","Level","NowExp","StartHP","StatusPoint","Elixir",
           "StartStrength","StartVital","StartAgility","StartMagic","StartSkill","StartSkill_Lv",
           "StartCritRate","StartCritDamage","StartAccuracy","StartAtkSpeed","StartPhysicalResist",
           "StartMagicResist","StartEvasion","StartRecoveryRate","StartGrowthRate","StartDropRate"]
def build_playerstatus():
    wb = Workbook(); ws = wb.active; ws.title = "PlayerStatusData"
    ws.append(COLS_PS); ws.append(["int","string"] + ["int"]*11 + ["multiline","multiline"] + ["float"]*10)
    ws.append(["ID","직업","숙련도","모험가등급","레벨","경험치","시작HP","능력점수","엘릭서",
               "근력","활력","민첩","마력","시작능력(줄바꿈)","시작능력Lv(줄바꿈)",
               "치명타%","치명피해%","명중%","공속%","물리저항%","마법저항%","회피%","회복%","성장률%","획득률%"])
    for e in j("PlayerStatusData.json")["playerStatuses"]:
        row = [cell(e.get(c)) for c in COLS_PS[:13]]
        row.append("\n".join(e.get("StartSkills") or []))
        row.append("\n".join(str(x) for x in (e.get("StartSkillLevels") or [])))
        row += [cell(e.get(c)) for c in COLS_PS[15:]]
        ws.append(row)
    save(wb, "PlayerStatusData.xlsx")

# ── CreatureData.xlsx (3행 헤더, 신규 로컬 파서 순서: GS 27컬럼 + 수기 JSON 전용 10컬럼) ──
COLS_CR = ["DataId","Name","PrefabLabel","IconLabel","MaxHp","PhysicalAtk","Def","MoveSpeed","TotalExp",
           "RewardGold","DropItemId","DropCount","DropChance","MagicAtk","CritRate","CritDamage","Accuracy",
           "AtkSpeed","PhysicalResist","MagicResist","Evasion","RecoveryRate","UniqueSkillName","UniqueSkillDesc",
           "AttackSfx","OnHitSfx","DieSfx",
           "AnimController","Habitat","DmgAmp","Rarity","MonsterAttackDelay","EncounterRate","InterestRate",
           "MaxKillCount","SpawnY","DamageFontYOffset"]
def build_creature():
    wb = Workbook(); ws = wb.active; ws.title = "CreatureData"
    ws.append(COLS_CR); ws.append([""]*len(COLS_CR))
    ws.append(["ID","이름(로컬키)","프리팹","아이콘","최대HP","물공","방어","이속","경험치","골드",
               "드랍ID","드랍수","드랍확률0~1","마공","치명%","치명피해%","명중%","공속%","물저%","마저%",
               "회피%","회복%","고유기술명(로컬키)","고유기술설명(로컬키)","공격SFX","피격SFX","사망SFX",
               "애니컨트롤러","서식지(로컬키)","피해증폭%","희귀도","공격대기(초)","조우확률0~1","관심확률0~1",
               "최대킬카운트","스폰Y","데미지폰트Y"])
    for e in j("CreatureData.json")["creatures"]:
        ws.append([cell(e.get(c)) for c in COLS_CR])
    save(wb, "CreatureData.xlsx")

# ── ItemData.xlsx (시트 3개: InventoryItem/RequisiteItem/Equipment, 1행 헤더, Grade=enum명) ──
def build_item():
    wb = Workbook(); wb.remove(wb.active)
    for sheet, src in [("InventoryItem","InventoryItem.json"),("RequisiteItem","RequisiteItem.json"),("Equipment","Equipment.json")]:
        ws = wb.create_sheet(sheet)
        ws.append(["DataId","Name","Grade","Description","SpriteName","FoundLocation"])
        for e in j(src)["items"]:
            g = e.get("Grade")
            ws.append([e["DataId"], e.get("Name",""), EQUIP_GRADE[g] if isinstance(g,int) else g,
                       e.get("Description",""), e.get("SpriteName",""), e.get("FoundLocation","")])
    save(wb, "ItemData.xlsx")

# ── LocalizationData.xlsx (단일 시트, 1행 헤더) ──
def build_localization():
    wb = Workbook(); ws = wb.active; ws.title = "Localization"
    ws.append(["TextKey","Kor","Eng"])
    data = j("LocalizationData.json"); root = list(data)[0]
    for e in data[root]:
        ws.append([e["TextKey"], e.get("Kor",""), e.get("Eng","")])
    save(wb, "LocalizationData.xlsx")

# ── DialogueData.xlsx (시트=그룹키; JSON 그룹 + 기존 파일의 미래용 GuildOnQuestText 보존) ──
def build_dialogue():
    keep = {}
    old = os.path.join(XL, "DialogueData.xlsx")
    if os.path.exists(old):
        owb = load_workbook(old)
        if "GuildOnQuestText" in owb.sheetnames:
            keep["GuildOnQuestText"] = [r for r in owb["GuildOnQuestText"].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    wb = Workbook(); wb.remove(wb.active)
    for g in j("DialogueScriptData.json")["groups"]:
        ws = wb.create_sheet(g["GroupKey"])
        ws.append(["TextKey","Kor","Eng"])
        for line in g["Lines"]:
            ws.append([line["TextKey"], line.get("Kor",""), line.get("Eng","")])
    for name, rows in keep.items():
        ws = wb.create_sheet(name); ws.append(["TextKey","Kor","Eng"])
        for r in rows: ws.append(list(r[:3]))
    save(wb, "DialogueData.xlsx")

# ── TravelData.xlsx (시트 TravelNode/TravelCondition, 1행 헤더 — ParseTravelSheets 순서) ──
def build_travel():
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("TravelNode")
    ws.append(["Id","ParentId","Category","Name","Sprite","AreaType","ClickBtnId","EnterBtnId"])
    for n in j("TravelNodeData.json")["nodes"]:
        ws.append([cell(n.get(c)) for c in ["Id","ParentId","Category","Name","Sprite","AreaType","ClickBtnId","EnterBtnId"]])
    ws = wb.create_sheet("TravelCondition")
    ws.append(["NodeId","Type","Target","Value","FailMessage"])
    for c in j("TravelConditionData.json")["conditions"]:
        ws.append([cell(c.get(k)) for k in ["NodeId","Type","Target","Value","FailMessage"]])
    save(wb, "TravelData.xlsx")

# ── BuffData.xlsx (3행 헤더 — ParseBuffDataSheet 순서) ──
def build_buff():
    wb = Workbook(); ws = wb.active; ws.title = "BuffData"
    ws.append(["DataId","Name","Description","IconLabel","Timed","Duration"])
    ws.append(["int","string","string","string","bool","float"])
    ws.append(["ID","이름(로컬키)","설명(로컬키)","아이콘","시간제여부","지속(초)"])
    for e in j("BuffData.json")["buffs"]:
        ws.append([cell(e.get(c)) for c in ["DataId","Name","Description","IconLabel","Timed","Duration"]])
    save(wb, "BuffData.xlsx")

if __name__ == "__main__":
    build_playerstatus(); build_creature(); build_item()
    build_localization(); build_dialogue(); build_travel(); build_buff()
    print("완료 — 이후 Unity: Tools/ParseAllLocal 실행 후 JsonData diff로 검증할 것")
