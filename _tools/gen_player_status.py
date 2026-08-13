# -*- coding: utf-8 -*-
"""PlayerStatusData.xlsx 를 레벨 1~300 스탯 테이블로 재생성.

배경:
  - 기존 시트는 데이터가 1행(DataId=99999999)뿐이고, MakeDict가 DataId로 키를 잡는 탓에
    PlayerController.RefreshStats의 PlayerStatusDic[레벨] 조회가 항상 실패했다.
    → 실제 스탯은 코드 fallback 공식이 담당하는 상태였다(코드 주석도 이를 인지).
  - 이 스크립트는 DataId=Level 로 300행을 채워 데이터 주도 경로를 살린다.

밸런스 설계(난이도/LOD):
  - Lv1~40 구간은 기존 fallback 공식과 **완전히 동일한 값**을 넣어 현재 체감을 보존한다.
      HP=100+20*(lv-1), 근력=15+3*(lv-1), 마력=10+2*(lv-1), 방어=1*(lv-1), 명중=90
  - Lv41부터는 성장 폭을 밴드별로 키운다. 요구경험치는 등비인데 스탯이 선형이면
    후반에 "레벨업해도 안 강해지는데 경험치만 폭증"하는 붕괴가 생기기 때문.

주의: 헤더 3줄(필드명/타입/한글설명)은 원본 셀을 그대로 두고 데이터 행만 교체한다.
실행: python _tools/gen_player_status.py  → 이후 Unity에서 Tools/ParseAllLocal
"""
import os
import openpyxl

XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                    "..", "Assets", "@Resources", "Data", "Excel", "PlayerStatusData.xlsx")

# (구간 끝 레벨, HP증가, 근력증가, 마력증가, 방어증가)
BANDS = [
    (40,   20,  3,  2, 1),   # 기존 fallback 공식과 동일 — 초반 체감 보존
    (100,  35,  5,  3, 2),
    (180,  60,  8,  5, 3),
    (250, 100, 13,  8, 5),
    (300, 160, 20, 12, 8),
]


def growth_for(level):
    for end, hp, st, mg, df in BANDS:
        if level <= end:
            return hp, st, mg, df
    return BANDS[-1][1:]


def build_rows():
    hp, st, mg, df = 100, 15, 10, 0          # Lv1 기준값 (fallback과 동일)
    out = []
    for lv in range(1, 301):
        if lv > 1:
            ghp, gst, gmg, gdf = growth_for(lv)
            hp += ghp
            st += gst
            mg += gmg
            df += gdf
        out.append(dict(lv=lv, hp=hp, st=st, mg=mg, df=df))
    return out


def main():
    path = os.path.normpath(XLSX)
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]

    # StartDef 컬럼 추가 (없을 때만) — 헤더 3줄 규칙에 맞춰 이름/타입/설명
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    if "StartDef" not in headers:
        col = ws.max_column + 1
        ws.cell(1, col, "StartDef")
        ws.cell(2, col, "int")
        ws.cell(3, col, "방어력")
    def_col = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)].index("StartDef") + 1

    idx = {name: i + 1 for i, name in enumerate(ws.cell(1, c).value for c in range(1, ws.max_column + 1))}

    # 기존 데이터 행(4행 이후) 제거 후 재작성
    if ws.max_row >= 4:
        ws.delete_rows(4, ws.max_row - 3)

    for r, row in enumerate(build_rows(), start=4):
        def put(name, value):
            if name in idx:
                ws.cell(r, idx[name], value)

        put("DataId", row["lv"])          # ← 레벨 키 조회가 되도록 DataId=Level
        put("Class", "Traveler")
        put("ClassProficiency", 0)
        put("AdvRank", 0)
        put("Level", row["lv"])
        put("NowExp", 0)
        put("StartHP", row["hp"])
        put("StatusPoint", 0)
        put("Elixir", 0)
        put("StartStrength", row["st"])
        put("StartVital", 0)
        put("StartAgility", 0)
        put("StartMagic", row["mg"])
        put("StartSkill", "0")
        put("StartSkill_Lv", "0")
        put("StartCritRate", 0)
        put("StartCritDamage", 150)
        put("StartAccuracy", 90)          # fallback과 동일(콤보 끊김 방지)
        put("StartAtkSpeed", 100)
        put("StartPhysicalResist", 0)
        put("StartMagicResist", 0)
        put("StartEvasion", 0)
        put("StartRecoveryRate", 0)
        put("StartGrowthRate", 100)
        put("StartDropRate", 100)
        ws.cell(r, def_col, row["df"])

    wb.save(path)
    print("saved:", path)
    rows = {r["lv"]: r for r in build_rows()}
    for lv in (1, 10, 40, 41, 100, 180, 250, 300):
        r = rows[lv]
        print(f"  Lv{lv:3d}  HP {r['hp']:6,}  ATK {r['st']:5,}  MAG {r['mg']:5,}  DEF {r['df']:4,}")


if __name__ == "__main__":
    main()
