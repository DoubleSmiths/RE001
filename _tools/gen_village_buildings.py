# -*- coding: utf-8 -*-
"""마을 건물 3종(수도원·대장간·상점)의 로컬라이즈 문자열 + 퀘스트 해금 행을 넣는다.

로컬라이즈 시트는 외부에서 재생성되면 코드가 추가한 키가 날아간 전례가 있어(loc_sheet_regenerated),
기존 키가 있으면 **덮어쓰지 않고 건너뛴다**.
"""
import io
import os

import openpyxl

BASE = os.path.abspath('.')
LOC = os.path.join(BASE, 'Assets', '@Resources', 'Data', 'Excel', 'LocalizationData.xlsx')
QST = os.path.join(BASE, 'Assets', '@Resources', 'Data', 'Excel', 'QuestData.xlsx')

# (TextKey, Kor, Eng)
ROWS = [
    # 수도원
    ('MonasteryUnlockText', '수도원에 출입할 수 있게 되었습니다.\n[축성]으로 직업 능력을 받으세요.',
     'The Monastery is now open to you.\nReceive your class blessings there.'),
    ('MonasteryNothingText', '준비 중', 'LOCKED'),

    # 대장간
    ('ForgeText', '대장간', 'FORGE'),
    ('ForgeUnlockText', '대장간이 문을 열었습니다.\n재료를 모아 장비를 만들 수 있습니다.',
     'The Forge has opened.\nGather materials to craft equipment.'),
    ('ForgeCraftBtnText', '제작', 'CRAFT'),
    ('ForgeEnhanceBtnText', '강화', 'ENHANCE'),
    ('ForgeCraftTitleText', '장비 제작', 'CRAFT EQUIPMENT'),
    ('ForgeNoCraftableText', '지금 만들 수 있는 장비가 없습니다.', 'Nothing can be crafted right now.'),
    ('ForgeCraftedText', '제작 완료', 'CRAFTED'),
    ('ForgeOwnedText', '보유 중', 'OWNED'),
    ('ForgeNoEquipText', '손질할 장비가 없습니다.', 'You have no equipment to work on.'),

    # 상점
    ('ShopText', '상점', 'SHOP'),
    ('ShopUnlockText', '상점이 문을 열었습니다.\n물약과 재료를 사고팔 수 있습니다.',
     'The Shop has opened.\nBuy and sell potions and materials.'),
    ('ShopBuyBtnText', '구매', 'BUY'),
    ('ShopSellBtnText', '판매', 'SELL'),
    ('ShopBuyTitleText', '판매 목록', 'FOR SALE'),
    ('ShopSellTitleText', '되팔기', 'SELL BACK'),
    ('ShopNotEnoughGoldText', '골드가 부족합니다.', 'Not enough gold.'),
    ('ShopNothingToSellText', '되팔 물건이 없습니다.', 'Nothing to sell.'),
    ('ShopLockedItemText', 'Lv.{0} 부터', 'From Lv.{0}'),
    ('ShopPurchasedText', '구매했습니다.', 'Purchased.'),
    ('ShopSoldText', '판매했습니다.', 'Sold.'),
    ('ShopSellPriceText', '판매가', 'SELL PRICE'),
    ('ShopOwnedCountText', '보유', 'OWNED'),
]

# (QuestID, UnlockType, TargetID) — 승급 보고 퀘스트마다 하나씩 열어 준다.
#  7 = 2급 승급 보고, 9 = 3급, 11 = 4급. 연금공방이 5번에서 열린 것과 같은 리듬.
UNLOCKS = [
    (7,  'UI_BUTTON', 'UNLOCK_FORGE_BUTTON'),
    (9,  'UI_BUTTON', 'UNLOCK_SHOP_BUTTON'),
    (11, 'UI_BUTTON', 'UNLOCK_MONASTERY_BUTTON'),
]


def add_localization():
    wb = openpyxl.load_workbook(LOC)
    ws = wb['Localization']
    have = set()
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row and row[0]:
            have.add(str(row[0]).strip())

    added, skipped = 0, []
    for key, kor, eng in ROWS:
        if key in have:
            skipped.append(key)
            continue
        ws.append([key, kor, eng])
        added += 1
    wb.save(LOC)
    print(f'로컬라이즈: {added}건 추가, 기존 유지 {len(skipped)}건 {skipped}')


def add_unlocks():
    wb = openpyxl.load_workbook(QST)
    ws = wb['Quest_Unlock']
    have = set()
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row and row[0] is not None:
            have.add((row[0], str(row[1]), str(row[2])))

    added = 0
    for qid, utype, target in UNLOCKS:
        if (qid, utype, target) in have:
            continue
        ws.append([qid, utype, target])
        added += 1
    wb.save(QST)
    print(f'퀘스트 해금: {added}건 추가')


if __name__ == '__main__':
    add_localization()
    add_unlocks()
    print('→ Unity에서 Tools/ParseAllLocal + Tools/GenerateLocalizeKeys 필요')
