"""첫 숲에서 제작하는 견습생 무기 두 종을 추가한다.

기본 dry-run, --apply로 엑셀 네 파일만 수정한다. 기존 행은 변경하지 않는다.
적용 후 Unity Tools/ParseAllLocal로 JSON을 재생성해야 한다.
"""
import argparse
from copy import copy
from pathlib import Path

import openpyxl

EXCEL = Path(__file__).resolve().parents[1] / 'Assets/@Resources/Data/Excel'
WEAPONS = (
    {'id': 30061, 'recipe': 40060, 'kor': '견습생의 장검', 'eng': 'Apprentice Longsword',
     'sprite': 'Sword 32 (5)', 'atk': 18, 'crit': 0, 'enhance_atk': 9,
     'materials': ((11001, 3), (10001, 3), (10002, 2))},
    {'id': 30062, 'recipe': 40061, 'kor': '견습생의 세검', 'eng': 'Apprentice Rapier',
     'sprite': 'Sword01', 'atk': 15, 'crit': 10, 'enhance_atk': 7,
     'materials': ((11001, 3), (10002, 3), (11004, 2))},
)
GUIDE_TEXT = {
    'GrowthClassGuideText': (
        '처치마다 숙련도 +1\n100이면 축성 가능',
        'Kills: +1 proficiency\nAnoint at 100'),
    'GrowthMaterialHintText': (
        '재료 이름을 누르면 획득 장소를 확인합니다.',
        'Tap a material to see where to find it.'),
    'GrowthItemRarityLabelText': ('등급:', 'Rarity:'),
    'GrowthItemFoundLabelText': ('발견장소:', 'Found:'),
}


def records(ws, start):
    headers = [cell.value for cell in ws[1]]
    result = {}
    for row in ws.iter_rows(min_row=start):
        key = row[0].value
        if key is None:
            continue
        assert key not in result, f'{ws.title} 중복 키: {key}'
        result[key] = dict(zip(headers, (cell.value for cell in row)))
    return result


def check_material_sources(excel):
    # ItemData는 메타 행 없이 2행부터, CreatureData는 4행부터 실제 데이터다.
    item_book = openpyxl.load_workbook(excel / 'ItemData.xlsx', read_only=True)
    creature_book = openpyxl.load_workbook(excel / 'CreatureData.xlsx', read_only=True)
    try:
        items = records(item_book.active, 2)
        creatures = records(creature_book.active, 4)
        habitat = creatures[2001]['Habitat']
        first_creatures = [row for row in creatures.values() if row['Habitat'] == habitat]
        for item_id in {item_id for weapon in WEAPONS for item_id, _ in weapon['materials']}:
            assert item_id in items, f'재료 ItemData 누락: {item_id}'
            sources = [row['DataId'] for row in first_creatures
                       if row['DropItemId'] == item_id and row['DropChance'] > 0 and row['DropCount'] > 0
                       and row['EncounterRate'] > 0]
            assert sources, f'첫 숲에 출현하는 마물의 실제 드롭이 아닌 재료: {item_id}'
            print(f'첫 숲 재료 {item_id}: 마물 {sources}')
    finally:
        item_book.close()
        creature_book.close()


def upgrade(apply=False, excel=EXCEL):
    excel = Path(excel)
    check_material_sources(excel)
    names = ('EquipmentData', 'EquipEnhanceData', 'ForgeRecipeData', 'LocalizationData')
    books = {name: openpyxl.load_workbook(excel / f'{name}.xlsx') for name in names}
    added = {name: 0 for name in names}
    try:
        existing = {name: records(book.active, 2 if name == 'LocalizationData' else 4)
                    for name, book in books.items()}
        original_rows = {name: book.active.max_row for name, book in books.items()}
        before = {name: [(cell.coordinate, cell.value, copy(cell._style))
                         for row in book.active for cell in row] for name, book in books.items()}

        def append(name, values):
            ws = books[name].active
            key_name = 'TextKey' if name == 'LocalizationData' else 'DataId'
            key = values[key_name]
            if key in existing[name]:
                assert all(existing[name][key].get(column) == value for column, value in values.items()), (
                    f'{name} 기존 ID/번역 충돌: {key}', existing[name][key])
                return
            header = {cell.value: cell.column for cell in ws[1]}
            assert set(values) <= set(header), f'{name} 필수 컬럼 누락'
            # ID가 달라도 동일 장비 레시피/강화 단계가 이미 있으면 중복 생성하지 않는다.
            if name == 'ForgeRecipeData':
                assert not any(row['EquipDataId'] == values['EquipDataId'] for row in existing[name].values()), values
            if name == 'EquipEnhanceData':
                assert not any((row['EquipDataId'], row['Level']) == (values['EquipDataId'], values['Level'])
                               for row in existing[name].values()), values
            target = ws.max_row + 1
            for column_name, column in header.items():
                cell = ws.cell(target, column, values.get(column_name))
                cell._style = copy(ws.cell(original_rows[name], column)._style)
            existing[name][key] = values
            added[name] += 1

        for index, weapon in enumerate(WEAPONS):
            name_key = f'Equip{weapon["id"]}NameText'
            append('EquipmentData', {
                'DataId': weapon['id'], 'Name': name_key, 'SlotType': 'Weapon', 'Grade': 'Basic',
                'SpriteName': weapon['sprite'], 'MaxDurability': 110, 'BaseAtk': weapon['atk'],
                'BaseCritRate': weapon['crit'], 'BaseMagicAtk': 0,
            })
            recipe = {'DataId': weapon['recipe'], 'EquipDataId': weapon['id'], 'ReqLevel': 5, 'Gold': 90}
            for slot, (item_id, count) in enumerate(weapon['materials'], 1):
                recipe[f'Mat{slot}Id'] = item_id
                recipe[f'Mat{slot}Count'] = count
            append('ForgeRecipeData', recipe)
            for level in range(1, 10):
                append('EquipEnhanceData', {
                    'DataId': 540 + index * 9 + level, 'EquipDataId': weapon['id'], 'Level': level,
                    'AtkBonus': weapon['enhance_atk'], 'CritBonus': 0, 'MagicAtkBonus': 0,
                    'RequiredMaterialId': 11001, 'RequiredMaterialCount': 3 * level,
                })
            append('LocalizationData', {'TextKey': name_key, 'Kor': weapon['kor'], 'Eng': weapon['eng']})
        for key, (kor, eng) in GUIDE_TEXT.items():
            append('LocalizationData', {'TextKey': key, 'Kor': kor, 'Eng': eng})

        # 저장 전에 원래 모든 셀의 값과 스타일 불변을 확인한다. 신규 ID만 행 끝에 추가한다.
        for name, cells in before.items():
            ws = books[name].active
            assert all(ws[coordinate].value == value and ws[coordinate]._style == style
                       for coordinate, value, style in cells), f'기존 행 변경 감지: {name}'
            assert ws.max_row == original_rows[name] + added[name], name
            print(f'{name}: {added[name]}개 행 추가, apply={apply}')
        if apply:
            for name in names:
                if added[name]:
                    books[name].save(excel / f'{name}.xlsx')
    finally:
        for book in books.values():
            book.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    upgrade(parser.parse_args().apply)
