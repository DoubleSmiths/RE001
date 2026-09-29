"""전투 능력바 상태 번역을 추가한다. 기본 dry-run, --apply로 엑셀만 수정한다.

적용 후 Unity Tools/ParseAllLocal로 JSON을 재생성해야 한다.
"""
import argparse
from copy import copy
from pathlib import Path

import openpyxl

EXCEL = Path(__file__).resolve().parents[1] / 'Assets/@Resources/Data/Excel/LocalizationData.xlsx'
TRANSLATIONS = {
    'CombatAbilityReadyText': ('눌러서 사용', 'Tap to use'),
    'CombatAbilityCooldownText': ('재사용 대기', 'Cooling down'),
    'CombatAbilityFullHpText': ('체력이 가득합니다', 'Health is full'),
    'CombatAbilityNoTargetText': ('적을 기다리는 중', 'Waiting for a target'),
    'CombatAbilityUnavailableText': ('지금은 사용할 수 없습니다', 'Unavailable'),
}


def upgrade(apply=False):
    book = openpyxl.load_workbook(EXCEL)
    try:
        ws = book.active
        assert [cell.value for cell in ws[1]] == ['TextKey', 'Kor', 'Eng']
        existing = {}
        for row in ws.iter_rows(min_row=2):
            key = row[0].value
            if key is None:
                continue
            assert key not in existing, f'중복 키: {key}'
            existing[key] = tuple(cell.value for cell in row[1:3])
        added = 0
        for key, translations in TRANSLATIONS.items():
            if key in existing:
                assert existing[key] == translations, f'기존 번역 충돌: {key}'
                continue
            previous_row = ws.max_row
            target_row = previous_row + 1
            for column, value in enumerate((key, *translations), 1):
                cell = ws.cell(target_row, column, value)
                cell._style = copy(ws.cell(previous_row, column)._style)
            ws.row_dimensions[target_row].height = ws.row_dimensions[previous_row].height
            added += 1
            print(f'{key}: {translations[0]} / {translations[1]}')
        print(f'LocalizationData: {added}개 행 추가, apply={apply}')
        if apply and added:
            book.save(EXCEL)
    finally:
        book.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    upgrade(parser.parse_args().apply)
