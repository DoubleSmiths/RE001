"""영문 용어 통일: 설명문의 'minion(s)' → 게임 고유명 'Minimon'.

  UI 라벨은 이미 MINIMON(도감·나의 미니몬 등)인데 능력 설명 약 500행은 'minion' 이라 한 화면에 두 이름이 섞였다.
  영문 제목도 'Minimon Quest' 다. 복수형은 고유명처럼 그대로 'Minimon'(스토어 문구와 같은 표기).
  대문자 라벨은 'MINIMON', 그 외는 'Minimon'. 키(TextKey)와 한국어는 건드리지 않는다.

  사용: python _tools/normalize_minimon_term.py           # 미리보기(건수 + 예시)
        python _tools/normalize_minimon_term.py --write   # LocalizationData.xlsx 저장 → Unity Tools/ParseAllLocal
"""
import re
import sys

import openpyxl

XLSX = "Assets/@Resources/Data/Excel/LocalizationData.xlsx"
PATTERN = re.compile(r"\b(minion)(s)?\b", re.IGNORECASE)


def replace(text):
    def sub(m):
        word = m.group(1)
        return "MINIMON" if word.isupper() else "Minimon"
    return PATTERN.sub(sub, text)


def main(write):
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Localization"]
    header = [c.value for c in ws[1]]
    eng = header.index("Eng")
    changed = []
    for row in ws.iter_rows(min_row=2):
        cell = row[eng]
        if not isinstance(cell.value, str) or not PATTERN.search(cell.value):
            continue
        new = replace(cell.value)
        if new != cell.value:
            changed.append((row[0].value, cell.value, new))
            cell.value = new
    for key, old, new in changed[:5]:
        print(f"{key}: {old[:70]!r} -> {new[:70]!r}")
    print(f"{len(changed)}행")
    if write and changed:
        wb.save(XLSX)
        print("저장 — Unity Tools/ParseAllLocal 로 JSON 재생성")


if __name__ == "__main__":
    main("--write" in sys.argv)
