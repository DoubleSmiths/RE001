# -*- coding: utf-8 -*-
"""v2 의상(치장) 8벌 생성기 (ADR0043) — CostumeData.xlsx.

  새 그림 없이 주인공 그림을 부위별로 다시 칠한다(CostumeManager). 한 벌 = 가죽·천 색(색상 이동·채도·밝기 배수)
  + 흰 옷·머리·강철에 곱할 색(Tint, 빈칸이면 그대로). 얼굴(살갗)은 어느 의상이든 그대로다.
  얻는 법: Default(처음부터) / Crystal(값) / Level(레벨) / Promotion(시련 단계) / Abyss(심연 최고 층).
  능력치는 없다(설계서 5장 치장 · 칭호와 같은 원칙) — 크리스탈 값은 치장 가격대(₩3,300~₩9,900 ≈ 300~900).

  값을 고칠 때는 여기를 고치고 다시 돌린 뒤 Unity Tools/ParseAllLocal.
    python _tools/gen_costumes.py            # 엑셀을 쓴다
    python _tools/gen_costumes.py --check    # 저장된 엑셀이 생성 결과와 같은지만 본다
"""
import argparse
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL = os.path.join(ROOT, 'Assets', '@Resources', 'Data', 'Excel')

HEADER = [('DataId', 'int', '번호(목록 순서)'), ('NameKey', 'string', '이름 로컬라이즈 키'),
          ('Hue', 'float', '가죽·천 색상 이동(도)'), ('Sat', 'float', '가죽·천 채도 배수'), ('Val', 'float', '가죽·천 밝기 배수'),
          ('Tint', 'string', '흰 옷·머리·강철에 곱할 색(RRGGBB, 빈칸 = 그대로)'),
          ('Unlock', 'string', 'Default / Crystal / Level / Promotion / Abyss'), ('UnlockValue', 'int', '값·레벨·단계·층')]

# 이름 키, 색상 이동, 채도, 밝기, 흰색 곱, 얻는 법, 값
COSTUMES = [
    ('CostumeBaseName',    0,   1.0,  1.0,  None,     'Default',   0),
    ('CostumeCrimsonName', -15, 1.25, 1.1,  None,     'Crystal',   300),
    ('CostumeForestName',  95,  0.9,  0.95, 'E6FFD9', 'Level',     50),
    ('CostumeFrostName',   185, 0.8,  1.15, 'CCEBFF', 'Crystal',   450),
    ('CostumeNightName',   240, 0.6,  0.6,  '6B6B80', 'Promotion', 3),
    ('CostumeSakuraName',  -40, 0.7,  1.35, 'FFD9E6', 'Crystal',   600),
    ('CostumeGoldName',    15,  1.3,  1.3,  'FFEDB3', 'Crystal',   900),
    ('CostumeAbyssName',   250, 1.1,  0.8,  'CCB3FF', 'Abyss',     30),
]


def rows():
    return [(i,) + c for i, c in enumerate(COSTUMES, 1)]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true')
    args = ap.parse_args()

    assert sum(1 for c in COSTUMES if c[5] == 'Default') == 1, '기본 의상은 한 벌'
    assert all(c[5] in ('Default', 'Crystal', 'Level', 'Promotion', 'Abyss') for c in COSTUMES)
    print(f"의상 {len(COSTUMES)}벌 · 크리스탈 값 합계 {sum(c[6] for c in COSTUMES if c[5] == 'Crystal')}")

    path = os.path.join(EXCEL, 'CostumeData.xlsx')
    if args.check:
        wb = openpyxl.load_workbook(path)
        same = [tuple(r) for r in wb.active.iter_rows(min_row=4, values_only=True)] == rows()
        print(('같음 ' if same else '다름 ') + 'CostumeData.xlsx')
        sys.exit(0 if same else 1)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'CostumeData'
    for i in range(3):
        ws.append([h[i] for h in HEADER])
    for r in rows():
        ws.append(list(r))
    wb.save(path)
    print('저장 CostumeData.xlsx')


if __name__ == '__main__':
    main()
