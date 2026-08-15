# -*- coding: utf-8 -*-
"""CreatureData.IconLabel → 종별 고유 도감 카드(Mon_{DataId})로 교체.

기존에는 503종이 원형 30종(Mon_BreezingWindForest_0 …)을 공유해 도감을 열면
같은 그림이 스무 개씩 늘어섰다. AnimController(인게임 애니)는 원형 공유가 맞으므로
**IconLabel만** 바꾼다 — 도감/미니몬창/잡기창이 쓰는 그림이 이 열이다.

카드 PNG가 실제로 있는 종만 바꾼다(생성이 아직 안 끝났으면 그 종은 원형을 유지).
"""
import os

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XL = os.path.join(BASE, "Assets", "@Resources", "Data", "Excel", "CreatureData.xlsx")
CARDS = os.path.join(BASE, "Assets", "@Resources", "Sprite", "Monster", "Card")


def main():
    have = {f[:-4] for f in os.listdir(CARDS)} if os.path.isdir(CARDS) else set()
    wb = openpyxl.load_workbook(XL)
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    id_col = head.index("DataId") + 1
    icon_col = head.index("IconLabel") + 1

    changed = skipped = 0
    for row in range(4, ws.max_row + 1):
        did = ws.cell(row=row, column=id_col).value
        if did is None:
            continue
        card = f"Mon_{did}"
        if card not in have:
            skipped += 1
            continue
        if ws.cell(row=row, column=icon_col).value != card:
            ws.cell(row=row, column=icon_col, value=card)
            changed += 1
    wb.save(XL)
    print(f"IconLabel 교체 {changed}행 / 카드 없어 유지 {skipped}행")


if __name__ == "__main__":
    main()
