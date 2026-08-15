# -*- coding: utf-8 -*-
"""MinMonData 전체 생성 — 몬스터 503종 전부를 잡을 수 있게 만든다.

## 왜
  MinMonData에 행이 **3개**뿐이었다. `MinMonManager.TryCatch`는 첫 줄에서
  `MinMonDic.TryGetValue` 실패 시 곧바로 false를 반환하므로, 나머지 500종은
  도감에 뜨든 관심 창고에 담기든 **절대 잡히지 않았다**. 잡기 로직 자체는
  멀쩡했고 데이터가 비어 있던 것이다.

## 난이도 사다리
  몬스터 이름이 "심도 수식어 + 원형"(어린→전설의, 10단계 × 50종)으로 이미
  정렬돼 있다. 이 사다리를 그대로 포획 난이도로 쓴다.
    포획 확률   0.30 → 0.12   (희귀 ×0.7, 보스 ×0.4)
    확정 천장    8회 → 26회    (희귀 +5, 보스 +12)
  동행 성능은 원본 몬스터 공격력에 비례(40%)시켜 깊이 들어갈수록 강해진다.

## 관심 확률도 같이 손본다
  전 종이 InterestRate 0.5였다. 창고가 2칸인데 두 번에 한 번 관심이 뜨면
  "창고 가득" 경고만 계속 나온다. 희귀도에 따라 0.15/0.25/0.40으로 낮춘다 —
  귀한 놈일수록 관심이 간다는 쪽이 읽기도 낫다.

손으로 저작된 2001·2002·2003은 값을 보존한다.

사용: python _tools/gen_minmon.py   → xlsx 갱신 후 Unity에서 Tools/ParseAllLocal
"""
import io
import json
import os

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XL = os.path.join(BASE, "Assets", "@Resources", "Data", "Excel")
JS = os.path.join(BASE, "Assets", "@Resources", "Data", "JsonData")

MODS = ["어린", "거친", "굶주린", "성난", "사나운",
        "오래된", "타락한", "심연의", "왕관을 쓴", "전설의"]

# 손으로 저작된 행 — 덮어쓰지 않는다.
AUTHORED = {2001, 2002, 2003}

RARITY_CATCH = {"Common": 1.0, "Rare": 0.7, "Boss": 0.4}
RARITY_CEIL = {"Common": 0, "Rare": 5, "Boss": 12}
RARITY_SPD = {"Common": 100.0, "Rare": 110.0, "Boss": 125.0}
RARITY_CRIT = {"Common": 0.0, "Rare": 5.0, "Boss": 10.0}
RARITY_INTEREST = {"Common": 0.15, "Rare": 0.25, "Boss": 0.40}


def load(name, key=None):
    with io.open(os.path.join(JS, name), encoding="utf-8-sig") as f:
        d = json.load(f)
    return d[key or list(d.keys())[0]]


def tier_of(kor_name):
    for i, m in enumerate(MODS):
        if kor_name.startswith(m + " "):
            return i
    return 0


def rows():
    loc = {r["TextKey"]: r["Kor"] for r in load("LocalizationData.json")}
    out = []
    for c in load("CreatureData.json", "creatures"):
        did = c["DataId"]
        t = tier_of(loc.get(c["Name"], ""))
        rar = c.get("Rarity", "Common")

        chance = round(max(0.03, (0.30 - t * 0.02) * RARITY_CATCH.get(rar, 1.0)), 3)
        ceil = 8 + t * 2 + RARITY_CEIL.get(rar, 0)

        out.append({
            "DataId": did,
            "Name": c["Name"],
            "IconLabel": f"Mon_{did}",             # 종별 고유 도감 카드
            "PrefabLabel": c.get("PrefabLabel", "Monster"),
            "PetAnimLabel": c.get("AnimController", ""),   # 동행도 제 모습으로 움직이게
            "MonsterDataId": did,
            "BaseAtk": round(c.get("PhysicalAtk", 0) * 0.4, 1),
            "BaseAtkSpeed": RARITY_SPD.get(rar, 100.0),
            "BaseCritRate": RARITY_CRIT.get(rar, 0.0),
            "AtkGrade": "E", "AtkSpeedGrade": "E", "CritGrade": "E",
            "MaxTameCount": ceil,
            "CatchChance": chance,
            "_interest": RARITY_INTEREST.get(rar, 0.15),
        })
    return out


def write_sheet(path, sheet, header_rows, records, cols):
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    ws.delete_rows(header_rows + 1, ws.max_row)
    for i, r in enumerate(records):
        for j, col in enumerate(cols):
            ws.cell(row=header_rows + 1 + i, column=j + 1, value=r[col])
    wb.save(path)


def main():
    recs = rows()

    # 저작 행 보존
    mm_path = os.path.join(XL, "MinMonData.xlsx")
    wb = openpyxl.load_workbook(mm_path)
    ws = wb["MinMonData"]
    cols = [c.value for c in ws[1]]
    keep = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row and row[0] in AUTHORED:
            keep[row[0]] = dict(zip(cols, row))
    wb.close()

    for r in recs:
        if r["DataId"] in keep:
            r.update({k: v for k, v in keep[r["DataId"]].items() if k in r})

    write_sheet(mm_path, "MinMonData", 3, recs, cols)
    print(f"MinMonData.xlsx — {len(recs)}행 (저작 보존 {len(keep)})")

    # CreatureData: 관심 확률만 조정. 다른 열은 건드리지 않는다.
    cd_path = os.path.join(XL, "CreatureData.xlsx")
    wb = openpyxl.load_workbook(cd_path)
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    ci = head.index("InterestRate") + 1
    by_id = {r["DataId"]: r for r in recs}
    id_col = head.index("DataId") + 1
    n = 0
    for row in range(4, ws.max_row + 1):
        did = ws.cell(row=row, column=id_col).value
        if did in by_id:
            ws.cell(row=row, column=ci, value=by_id[did]["_interest"])
            n += 1
    wb.save(cd_path)
    print(f"CreatureData.xlsx — InterestRate {n}행 조정")

    lo = min(recs, key=lambda r: r["DataId"])
    hi = max(recs, key=lambda r: r["DataId"])
    print(f"  예: {lo['DataId']} 확률 {lo['CatchChance']} 천장 {lo['MaxTameCount']} 공격 {lo['BaseAtk']}")
    print(f"      {hi['DataId']} 확률 {hi['CatchChance']} 천장 {hi['MaxTameCount']} 공격 {hi['BaseAtk']}")


if __name__ == "__main__":
    main()
