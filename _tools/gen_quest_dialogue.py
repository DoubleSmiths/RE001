# -*- coding: utf-8 -*-
"""추가한 메인 퀘스트(8~55)에 접수원 대사를 붙인다.

## 왜
  1~7번은 **대화로 임무를 받고 대화로 보고**한다. 그런데 새로 이어 붙인 8~55번은
  DialogueID가 비어 있어서, 협회에 가도 잡담(GuildNormalText)만 나오고 임무가 뭔지
  알려주는 사람이 없었다. "이전 퀘스트는 대화를 하고 받는데 추가한 건 그런 게 없다"는 지적.

## 기존 방식(그대로 따른다)
  · 대화 그룹 = DialogueData.xlsx의 **시트 하나**. 시트명이 곧 GroupKey.
    열은 TextKey / Kor / Eng. QuestData.DialogueID가 이 GroupKey를 가리킨다.
  · 대사를 다는 쪽은 **보고 퀘스트**(BTN_CLICK BTN_GUILD_QUEST)다.
    그 대화가 "이번 것 확인 → 승급 → **다음에 뭘 가져와야 하는지**"까지 말해 준다.
    수집 퀘스트(ITEM_COUNT)에는 대사가 없다 — 3번·6번이 그랬다.
  · 그래서 7번 대사(GuildQuest3Text)의 "협회에서 준비 중입니다" 문장도 손봐야 한다.
    거기서 8번 임무를 주지 않으면 새 체인의 첫 단추가 비어 버린다.

  화법: 접수원 아가씨 — 사무적인 존댓말, 색 태그는 원문 관례를 따른다
    재료 `[<#2DFF00><b>이름</b></color>] <#2DFF00><b>수</b></color>개`
    지역 `[<#B0FF06>이름</color>]`   등급 `<#00B9FF>[<#858585>n</color>]급 모험가</color>`

사용: python _tools/gen_quest_dialogue.py   → 이후 Unity에서 Tools/ParseAllLocal
"""
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_quest_chain import BASE, build_plan, item_names   # noqa: E402

DLG = os.path.join(BASE, "Assets", "@Resources", "Data", "Excel", "DialogueData.xlsx")
QST = os.path.join(BASE, "Assets", "@Resources", "Data", "Excel", "QuestData.xlsx")

FIRST_RANK = 3          # 1급=퀘4, 2급=퀘7까지 손으로 만들어져 있다. 새 체인은 3급부터.
GROUP_PREFIX = "GuildQuest"
FIRST_GROUP_NO = 4      # GuildQuest1~3Text가 이미 있다


def mat(name, count):
    return f"[<#2DFF00><b>{name}</b></color>] <#2DFF00><b>{count}</b></color>개"


def mat_en(name, count):
    return f"<color=#2DFF00><b>{count}</b></color> <color=#2DFF00><b>{name}</b></color>"


def place(name):
    return f"[<#B0FF06>{name}</color>]"


def rank(n):
    return f"<#00B9FF>[<#858585>{n}</color>]급 모험가</color>"


def rank_en(n):
    return f"<color=#00B9FF>[<color=#858585>{n}</color>]-rank Adventurer</color>"


def lines_for(step, nxt, rk, names):
    """보고 대화 5줄. nxt가 None이면 마지막 임무."""
    cur_mat = names.get(step["item"], str(step["item"]))
    out = [
        ("다시 오셨군요,\n" + mat(cur_mat, step["count"]) + "는 구해오셨나요?",
         "Back again.\nDid you bring " + mat_en(cur_mat, step["count"]) + "?"),
        ("….\n확인 되었습니다.\n협회 규정에 따라\n" + rank(rk) + " 승급을 진행하겠습니다.",
         "...Confirmed.\nIn accordance with guild regulations,\n"
         "your promotion to a " + rank_en(rk) + " will now proceed."),
    ]

    if nxt is not None:
        nxt_mat = names.get(nxt["item"], str(nxt["item"]))
        out.append((
            f"<#00B9FF>[<#858585>{rk}</color>]급 통행증</color>을 새로 발급해 드렸습니다.\n"
            f"이제 {place(nxt['kor'])}까지 출입하실 수 있습니다.\n"
            f"권장 수준은 <#FFFF00>{nxt['level']}</color> 이상입니다.",
            f"A new <color=#00B9FF>Rank {rk} Permit</color> has been issued.\n"
            f"You may now enter <color=#B0FF06>{nxt['kor']}</color>.\n"
            f"Recommended level: <color=#FFFF00>{nxt['level']}</color> or higher."))
        out.append((
            f"다음 <#00B9FF>[<#858585>{rk + 1}</color>]급</color>으로 진급하시려면\n"
            f"{mat(nxt_mat, nxt['count'])}가 필요합니다.\n"
            f"{place(nxt['kor'])}에 서식하는 마물에게서 얻으실 수 있습니다.",
            f"To reach <color=#00B9FF>rank {rk + 1}</color>,\n"
            f"you will need {mat_en(nxt_mat, nxt['count'])}.\n"
            f"They drop from the creatures of <color=#B0FF06>{nxt['kor']}</color>."))
        out.append(("그럼 다녀오신 뒤에 다시 뵙죠.", "Then I’ll see you when you return."))
    else:
        out.append((
            f"{rank(rk)} — 협회가 인정하는 최고 등급이십니다.\n"
            "모든 지역의 통행이 허가되었습니다.",
            f"{rank_en(rk)} — the highest rank the guild recognises.\n"
            "Passage to every region has been granted."))
        out.append((
            "여기까지가 협회가 준비한 마지막 임무입니다.\n그동안 수고 많으셨습니다.",
            "That was the last commission the guild had prepared.\nThank you for your service."))
        out.append(("언제든 다시 들러 주세요.", "Do drop by anytime."))
    return out


def write_group(wb, key, rows):
    if key in wb.sheetnames:
        del wb[key]
    ws = wb.create_sheet(key)
    ws.append(["TextKey", "Kor", "Eng"])
    for i, (kor, eng) in enumerate(rows, 1):
        ws.append([i, kor, eng])


def patch_handoff(wb, first, names):
    """7번 대사의 '준비 중입니다'를 실제 임무 안내로 바꾼다 — 새 체인의 첫 단추."""
    ws = wb["GuildQuest3Text"]
    m = names.get(first["item"], str(first["item"]))
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value in (4, "4"):
            ws.cell(row=row, column=2, value=(
                f"{rank(2)}로 승급되신 것을 축하드립니다.\n"
                f"이어서 {place(first['kor'])}로의 출입이 허가되었으며,\n"
                f"다음 <#00B9FF>[<#858585>{FIRST_RANK}</color>]급</color>으로 진급하시려면\n"
                f"{mat(m, first['count'])}가 필요합니다."))
            ws.cell(row=row, column=3, value=(
                f"Congratulations on your promotion to a {rank_en(2)}.\n"
                f"Access to <color=#B0FF06>{first['kor']}</color> has been granted.\n"
                f"To reach <color=#00B9FF>rank {FIRST_RANK}</color>, "
                f"you will need {mat_en(m, first['count'])}."))
            return True
    return False


def main():
    plan = build_plan()
    names = item_names()

    wb = openpyxl.load_workbook(DLG)
    made = []
    for i, step in enumerate(plan):
        nxt = plan[i + 1] if i + 1 < len(plan) else None
        key = f"{GROUP_PREFIX}{FIRST_GROUP_NO + i}Text"
        write_group(wb, key, lines_for(step, nxt, FIRST_RANK + i, names))
        made.append((step["report_id"], key))

    ok = patch_handoff(wb, plan[0], names)
    wb.save(DLG)

    # 보고 퀘스트 행에 DialogueID 기입
    qb = openpyxl.load_workbook(QST)
    ws = qb["Quest"]
    head = [c.value for c in ws[1]]
    id_col = head.index("QuestID") + 1
    dlg_col = head.index("DialogueID") + 1
    by_id = {qid: key for qid, key in made}
    n = 0
    for row in range(2, ws.max_row + 1):
        qid = ws.cell(row=row, column=id_col).value
        if qid in by_id:
            ws.cell(row=row, column=dlg_col, value=by_id[qid])
            n += 1
    qb.save(QST)

    print(f"대화 그룹 {len(made)}개 생성({made[0][1]}~{made[-1][1]}) / "
          f"보고 퀘스트 {n}건에 DialogueID 기입 / 7번 인계 문구 교체={ok}")
    print("→ Unity에서 Tools/ParseAllLocal 필요")


if __name__ == "__main__":
    main()
