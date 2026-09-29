# -*- coding: utf-8 -*-
"""v2 유물 30종 표 생성기 (ADR0035, 설계서 4.10).

  RelicData.xlsx — DataId, 이름 키, 등급(Common/Rare/Epic/Legendary), 효과(StatType), 1레벨 값, 레벨당 값, 아이콘.
  효과: Atk/Hp/Def = %, CritDmg = 치명 피해 %p, Gold/Exp = 획득 %, Drop = 획득률 %p, Interest = 관심 확률 %(곱).
  값(레벨 L) = Base + PerLevel × (L − 1), 최대 레벨은 BalanceConfig RELIC_MAX_LEVEL.
  아이콘은 _tools/aether_icons.py relic (파일명 Relic{DataId}) — RELIC_LOOK 과 순서·DataId 가 같다.

    python _tools/gen_relics.py            # 엑셀을 쓴다 → Unity Tools/ParseAllLocal
    python _tools/gen_relics.py --check    # 저장된 엑셀이 생성 결과와 같은지만 본다
    python _tools/gen_relics.py --loc      # fix_english_localization.py 에 붙일 이름 줄을 출력한다
"""
import argparse
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL = os.path.join(ROOT, 'Assets', '@Resources', 'Data', 'Excel', 'RelicData.xlsx')

# 등급별 (1레벨 값, 레벨당 값). 치명 피해·골드·경험치·획득률은 기본 %가 작게 느껴져 두 배.
VALUE = {
    'Common':    (1.0, 0.5),
    'Rare':      (2.0, 1.0),
    'Epic':      (4.0, 2.0),
    'Legendary': (8.0, 4.0),
}
DOUBLE = {'CritDmg', 'Gold', 'Exp', 'Drop'}
INTEREST = {'Common': (3.0, 1.5)}

# (DataId, 등급, 효과, 한국어 이름, 영어 이름) — 일반 12 · 희귀 9 · 영웅 6 · 전설 3
RELICS = [
    (8001, 'Common',    'Atk',      '녹슨 칼자루',     'Rusty Hilt'),
    (8002, 'Common',    'Atk',      '늑대 송곳니 부적', 'Wolf Fang Charm'),
    (8003, 'Rare',      'Atk',      '전사의 문장',     "Warrior's Crest"),
    (8004, 'Rare',      'Atk',      '불꽃 룬석',       'Flame Runestone'),
    (8005, 'Epic',      'Atk',      '용사의 건틀릿',   "Hero's Gauntlet"),
    (8006, 'Legendary', 'Atk',      '태양의 왕관',     'Crown of the Sun'),
    (8007, 'Common',    'Hp',       '이끼 낀 조약돌',   'Mossy Pebble'),
    (8008, 'Common',    'Hp',       '산딸기 병',       'Berry Jar'),
    (8009, 'Rare',      'Hp',       '생명의 도토리',   'Acorn of Life'),
    (8010, 'Epic',      'Hp',       '거인의 심장',     "Giant's Heart"),
    (8011, 'Legendary', 'Hp',       '세계수의 잎',     'World Tree Leaf'),
    (8012, 'Common',    'Def',      '방패 조각',       'Shield Shard'),
    (8013, 'Common',    'Def',      '거북 등딱지',     'Turtle Shell'),
    (8014, 'Rare',      'Def',      '철갑 견장',       'Iron Pauldron'),
    (8015, 'Epic',      'Def',      '수호자의 성배',   "Guardian's Chalice"),
    (8016, 'Common',    'CritDmg',  '날 선 부싯돌',     'Sharp Flint'),
    (8017, 'Rare',      'CritDmg',  '매의 눈',         "Hawk's Eye"),
    (8018, 'Epic',      'CritDmg',  '흡혈 송곳니',     'Vampire Fang'),
    (8019, 'Legendary', 'CritDmg',  '파멸의 단검',     'Dagger of Ruin'),
    (8020, 'Common',    'Gold',     '구리 동전',       'Copper Coin'),
    (8021, 'Common',    'Gold',     '가죽 지갑',       'Leather Purse'),
    (8022, 'Rare',      'Gold',     '행운의 편자',     'Lucky Horseshoe'),
    (8023, 'Epic',      'Gold',     '황금 두꺼비',     'Golden Toad'),
    (8024, 'Common',    'Exp',      '낡은 일지',       'Old Journal'),
    (8025, 'Rare',      'Exp',      '현자의 안경',     "Sage's Spectacles"),
    (8026, 'Rare',      'Exp',      '푸른 모래시계',   'Blue Hourglass'),
    (8027, 'Epic',      'Exp',      '별의 두루마리',   'Star Scroll'),
    (8028, 'Common',    'Drop',     '찢어진 자루',     'Torn Sack'),
    (8029, 'Rare',      'Drop',     '수집가의 갈고리', "Collector's Hook"),
    (8030, 'Common',    'Interest', '달콤한 꿀단지',   'Honey Pot'),
]

HEADER = [
    ('DataId', 'int', '번호'), ('NameKey', 'string', '이름 로컬라이즈 키'),
    ('Grade', 'string', '등급(Common/Rare/Epic/Legendary)'), ('StatType', 'string', '효과(Atk/Hp/Def/CritDmg/Gold/Exp/Drop/Interest)'),
    ('Base', 'float', '1레벨 값'), ('PerLevel', 'float', '레벨당 값'), ('IconLabel', 'string', '아이콘(어드레서블 주소)'),
]


def rows():
    out = []
    for did, grade, stat, _, _ in RELICS:
        base, per = INTEREST[grade] if stat == 'Interest' else VALUE[grade]
        if stat in DOUBLE:
            base, per = base * 2, per * 2
        out.append((did, f'Relic{did}NameText', grade, stat, base, per, f'Relic{did}'))
    return out


def loc_lines():
    return [f"    ('Relic{did}NameText', '{kor}', \"{eng}\")," for did, _, _, kor, eng in RELICS]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true')
    ap.add_argument('--loc', action='store_true')
    args = ap.parse_args()
    data = rows()
    counts = {g: sum(1 for r in RELICS if r[1] == g) for g in VALUE}
    assert counts == {'Common': 12, 'Rare': 9, 'Epic': 6, 'Legendary': 3}, counts
    assert len({r[0] for r in RELICS}) == len(RELICS)

    if args.loc:
        print('\n'.join(loc_lines()))
        return
    if args.check:
        wb = openpyxl.load_workbook(EXCEL)
        saved = [tuple(c for c in row) for row in wb.active.iter_rows(min_row=4, values_only=True)]
        same = saved == [tuple(r) for r in data]
        print(('같음 ' if same else '다름 ') + os.path.basename(EXCEL))
        sys.exit(0 if same else 1)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'RelicData'
    for i in range(3):
        ws.append([h[i] for h in HEADER])
    for r in data:
        ws.append(list(r))
    wb.save(EXCEL)
    print(f'저장 RelicData.xlsx ({len(data)}행) — 일반 {counts["Common"]} · 희귀 {counts["Rare"]} · 영웅 {counts["Epic"]} · 전설 {counts["Legendary"]}')


if __name__ == '__main__':
    main()
