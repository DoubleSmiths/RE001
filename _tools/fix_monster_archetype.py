# -*- coding: utf-8 -*-
"""몬스터의 **이름**을 기준으로 아이콘 스프라이트·드롭 재료·애니메이터를 다시 묶는다.

무엇이 틀어져 있었나
  gen_monsters 는 이름을 bases[k % 10] 으로 뽑고, 아이콘은 ANIMS[(did + k) % n] 으로 뽑았다.
  두 인덱스가 서로 무관해서 "슬라임인데 그림은 늑대"가 됐다.
  gen_world 는 거기에 더해 IconLabel 을 Mon_{지역}_{did % 10} 으로 다시 덮어썼다 — 역시 무관한 인덱스.
  드롭 재료도 같은 인덱스를 쓰므로 함께 어긋났다.

무엇을 기준으로 고치나
  **이름**이다. 사람이 보는 건 이름이고, 한글 이름(어린 슬라임 / 성난 숲거미 …)은 이미
  아키타입 30종과 1:1로 맞게 생성돼 있다. 이름에서 수식어를 떼어낸 base 로 아키타입을 찾고
  아이콘·재료·애니메이터를 전부 그 아키타입으로 맞춘다.
  → 이름과 그림이 어긋날 수 없는 구조가 된다.

사용:  python _tools/fix_monster_archetype.py [--dry]
이후 Unity 에서 Tools/ParseAllLocal 로 JSON 재생성 필요.
"""
import io
import os
import sys

import openpyxl

BASE = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
EXCEL = os.path.join(BASE, "Assets", "@Resources", "Data", "Excel")

MAT_ID0 = 11001          # 재료 아이템 첫 DataId (아키타입 30종과 1:1)

# 아키타입 계열 순서 = 재료 id 블록 순서. aether_icons.MONSTER_ARCHETYPES 와 같아야 한다.
FAMILIES = ["BreezingWindForest", "WhisperingCave", "RuinedSanctum"]

# 한글 base 이름 → (계열 index, 아키타입 index). gen_monsters.REGIONS[*]["bases"] 순서 그대로.
BASES = [
    # BreezingWindForest
    "슬라임", "숲거미", "들늑대", "독버섯", "바람정령", "멧돼지", "까마귀", "덩굴손", "여우", "말벌",
    # WhisperingCave
    "동굴박쥐", "두더지", "광부망령", "수정골렘", "지네", "이끼괴물", "석상병", "그림자", "도롱뇽", "광충",
    # RuinedSanctum
    "망령", "석상기사", "타락사제", "키메라", "가고일", "수호자", "환영", "타락정령", "제단수호물", "공허체",
]
BASE_TO_ARCH = {b: (i // 10, i % 10) for i, b in enumerate(BASES)}

# 이름 앞에 붙는 심도 수식어 (gen_monsters.MODS). base 를 찾으려면 떼어내야 한다.
MODS = ["어린", "거친", "굶주린", "성난", "사나운", "오래된",
        "타락한", "심연의", "왕관을 쓴", "전설의"]

# 손으로 만든 3종은 전용 아트가 있거나 이름이 특수하다 — 표에서 직접 지정한다.
#  2003 고블린은 goblin_monster.png + Goblin.controller 라는 **자체 아트**가 있어 아키타입에 넣지 않는다.
MANUAL = {
    2001: ("BreezingWindForest", 0),    # 슬라임      → slime blob
    2002: ("BreezingWindForest", 0),    # 라비 슬라임 → 같은 슬라임 아트 (색 변주)
}
# 전용 아트를 쓰는 몬스터 — 아키타입 대신 이 주소를 그대로 쓴다.
#  고블린은 goblin_monster.png + Goblin.controller 가 따로 있다. 아키타입에 억지로 끼우면
#  이름은 고블린인데 그림은 독버섯이 된다(실제로 그 상태였다).
OWN_ART = {
    # (IconLabel, AnimController)
    #  ⚠️ AnimController를 Mon_2003(레거시 개별 컨트롤러)로 두면 안 된다. 그 컨트롤러의 클립 5종은
    #     전부 **키프레임 1개**이고 그마저 Mon_BreezingWindForest_3(엉뚱한 아키타입 그림)을 가리켜,
    #     고블린이 전혀 움직이지 않는 정지 이미지로 나왔다.
    #     Goblin.controller에는 실제 다중 프레임 클립(공격 9키·사망 10키)과 진짜 고블린 아트가 있다.
    #     Run 상태만 없는데, 그건 CrossFade가 무시되고 직전 클립이 유지되므로 깨지지 않는다.
    2003: ("goblin_monster_0.sprite", "Goblin"),
}

# ⚠️ 메인 퀘스트가 요구하는 드롭. 아키타입 재료(11001~11030)로 덮으면 **게임이 진행 불가**가 된다.
#  퀘3 = 젤리 퐁(10001) 2개 / 퀘6 = 폴짝 젤리(10002) 5개인데, 드롭 집합이 11001~11030뿐이면
#  어떤 몬스터도 이 둘을 떨구지 않아 퀘3에서 영구 정지한다(실제로 그 상태였다).
#  아이템 스프라이트(Slime_Green_Drop / item_rabijelly)가 가리키는 대로 슬라임·라비 슬라임에 둔다.
#  Quest_Notify도 퀘3에서 산들바람 숲 입구로 유도하므로 설계 의도와 일치한다.
QUEST_DROPS = {
    2001: 10001,   # 슬라임      → 젤리 퐁
    2002: 10002,   # 라비 슬라임 → 폴짝 젤리
}


def load_korean_names():
    """LocalizationData.xlsx: TextKey → 한글 이름"""
    wb = openpyxl.load_workbook(os.path.join(EXCEL, "LocalizationData.xlsx"), read_only=True)
    ws = wb[wb.sheetnames[0]]
    out = {}
    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        if row[0]:
            out[str(row[0])] = str(row[1]) if row[1] else ""
    wb.close()
    return out


def base_of(korean):
    """'성난 숲거미' → '숲거미'. 수식어가 없으면 그대로."""
    s = korean.strip()
    for m in MODS:
        if s.startswith(m + " "):
            return s[len(m) + 1:].strip()
    return s


def main():
    dry = "--dry" in sys.argv
    names = load_korean_names()

    path = os.path.join(EXCEL, "CreatureData.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}

    fixed = same = unknown = 0
    unknown_names = set()

    for r in range(4, ws.max_row + 1):
        did = ws.cell(r, h["DataId"]).value
        if did is None:
            continue
        did = int(did)

        if did in OWN_ART:
            icon, anim = OWN_ART[did]
            if ws.cell(r, h["IconLabel"]).value == icon:
                same += 1
                continue
            if not dry:
                ws.cell(r, h["IconLabel"], icon)
                if "AnimController" in h:
                    ws.cell(r, h["AnimController"], anim)
            fixed += 1
            continue

        if did in MANUAL:
            fam_name, idx = MANUAL[did]
            fam = FAMILIES.index(fam_name)
        else:
            key = ws.cell(r, h["Name"]).value
            kor = names.get(str(key), "")
            b = base_of(kor)
            if b not in BASE_TO_ARCH:
                unknown += 1
                unknown_names.add("%s(%s)" % (kor, key))
                continue
            fam, idx = BASE_TO_ARCH[b]
            fam_name = FAMILIES[fam]

        icon = "Mon_%s_%d" % (fam_name, idx)
        drop = QUEST_DROPS.get(did, MAT_ID0 + fam * 10 + idx)

        cur_icon = ws.cell(r, h["IconLabel"]).value
        cur_drop = ws.cell(r, h["DropItemId"]).value if "DropItemId" in h else None
        cur_anim = ws.cell(r, h["AnimController"]).value if "AnimController" in h else icon

        # ⚠️ AnimController 도 함께 봐야 한다. 아이콘·드롭만 비교하면
        #    "그림은 맞는데 애니메이터만 옛 이름(Mon_2220)"인 행을 통째로 건너뛴다.
        #    실제로 17행이 그렇게 남아 전 구간 검증에서 잡혔다.
        if cur_icon == icon and cur_drop == drop and cur_anim == icon:
            same += 1
            continue

        if not dry:
            ws.cell(r, h["IconLabel"], icon)
            if "DropItemId" in h:
                ws.cell(r, h["DropItemId"], drop)
            # 애니메이터도 아키타입 단위로 (Task: 몬스터 4종 애니메이션)
            if "AnimController" in h:
                ws.cell(r, h["AnimController"], icon)
        fixed += 1

    if not dry:
        wb.save(path)

    print("아키타입 재정렬: 수정 %d / 이미맞음 %d / 이름불명 %d%s"
          % (fixed, same, unknown, "  (dry-run)" if dry else ""))
    if unknown_names:
        print("  이름을 아키타입에 매칭 못한 항목:", list(sorted(unknown_names))[:10])


if __name__ == "__main__":
    main()
