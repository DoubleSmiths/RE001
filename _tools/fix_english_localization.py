"""영문 로컬라이즈 마무리 — 기본 dry-run, --write 로만 저장, --check 로 결과 검증.

1. LocalizationData: 영문이 한글 그대로였던 77행(능력 이름/설명/이야기, SPDPotion2)의 Eng 만 채운다.
2. DialogueData: 영문 대사에 한글로 박혀 있던 재료·지역 이름을 LocalizationData 의 영문명으로 치환한다.
3. 런타임 코드의 하드코딩 한글을 걷어내기 위한 신규 키(Kor=기존 리터럴, Eng=번역)를 시트 끝에 추가하고,
   Define.LocalizeKey 에 상수를 동기화한다(#region LocalizeKey 안, 마지막 상수 뒤).
   (LITERAL_KEYS 는 코드가 문자열 리터럴로 부르는 키라 시트에만 넣고 Define 상수는 만들지 않는다.)

이후 Unity Tools/ParseAllLocal 로 JSON 을 재생성해야 한다.

    python _tools/fix_english_localization.py            # 미리보기
    python _tools/fix_english_localization.py --write    # 엑셀 2개 + Define.cs 저장
    python _tools/fix_english_localization.py --check    # 저장된 결과 재검증
"""
import argparse
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
EXCEL = ROOT / 'Assets/@Resources/Data/Excel'
LOC_PATH = EXCEL / 'LocalizationData.xlsx'
DLG_PATH = EXCEL / 'DialogueData.xlsx'
SCRIPTS = ROOT / 'Assets/@Scripts'
DEFINE = SCRIPTS / 'Utils/Define.cs'
# --check 가 중괄호 균형/LocalizeKey 참조를 검사하는 C# (이 작업이 손대는 파일)
CS_FILES = [
    'UI/Popup/UI_SkillDetailPopup.cs', 'UI/SubItem/UI_Village1BtnGroup.cs', 'UI/Popup/UI_BestiaryPopup.cs',
    'UI/Popup/UI_TamePopup.cs', 'UI/Popup/UI_AttactedPopup.cs', 'Managers/Contents/AbilityManager.cs',
    'UI/Scene/UI_GameScene.cs', 'Utils/Define.cs',
]

HANGUL = re.compile('[ᄀ-ᇿ㄰-㆏ꥠ-꥿가-힣ힰ-퟿]')
HANGUL_RUN = re.compile('[ᄀ-ᇿ㄰-㆏가-힣]+(?: [ᄀ-ᇿ㄰-㆏가-힣]+)*')

# ─────────────────────────────────────────────────────────────────────────────
# 1) 영문이 한글 그대로였던 행 — Eng 만 채운다 (Kor 은 그대로)
# ─────────────────────────────────────────────────────────────────────────────
SKILL_ENG = {
    # ── 능력 이름 ──
    'Skill5101NameText': 'Flurry Habit',
    'Skill5102NameText': 'Lingering Focus',
    'Skill5103NameText': 'Impact Sense',
    'Skill5104NameText': 'Critical Flow',
    'Skill5106NameText': 'Lesson of a Miss',
    'Skill5107NameText': 'Insight',
    'Skill5108NameText': 'Deflection',
    'Skill5110NameText': 'Toughened Skin',
    'Skill5111NameText': 'On the Brink',
    'Skill5112NameText': 'Coup de Grace',
    'Skill5114NameText': 'Healing Afterglow',
    'Skill5115NameText': 'Combat Breathing',
    'Skill5117NameText': 'Focus',
    'Skill5118NameText': 'Acceleration',
    'Skill5119NameText': 'Iron Wall',
    'Skill5120NameText': 'Keenness',
    'Skill5122NameText': 'Shatter',
    'Skill5123NameText': 'Brawn',
    'Skill5124NameText': 'Mana Circuit',
    'Skill5126NameText': 'Tempering',
    'Skill5127NameText': 'Weak Point Sense',
    'Skill5128NameText': 'Killing Blow',
    'Skill5129NameText': 'Precision',
    'Skill5130NameText': 'Swift Assault',
    'Skill5132NameText': 'Healing Affinity',
    'Skill5133NameText': 'Vessel of Memory',
    'Skill5134NameText': 'Endurance',
    'Skill5135NameText': 'Curiosity',
    'Skill5136NameText': "Merchant's Knack",
    'Skill5137NameText': 'Luck',
    'Skill5138NameText': 'Rapport',
    'Skill5140NameText': 'Scent of Affinity',
    'Skill5141NameText': 'Herbal Lore',
    'Skill5143NameText': "Blacksmith's Eye",
    'Skill5144NameText': 'Quick Steps',
    # ── 능력 설명 ──
    'Skill5101DescText': 'A 25% chance to deal bonus damage when attacking.',
    'Skill5102DescText': 'A 30% chance for your next strike to become sharper after attacking.',
    'Skill5103DescText': 'A 20% chance to restore a little HP when landing a hit.',
    'Skill5104DescText': 'A 45% chance to increase attack speed on a critical hit.',
    'Skill5106DescText': 'A 70% chance to increase accuracy when an attack misses.',
    'Skill5107DescText': 'A 50% chance to increase critical rate when the enemy dodges.',
    'Skill5108DescText': 'A 55% chance to deal counter damage after a successful dodge.',
    'Skill5110DescText': 'A 35% chance to harden your defense when hit.',
    'Skill5111DescText': 'Greatly increases physical attack when HP falls to 30% or below.',
    'Skill5112DescText': "Deals bonus damage when the enemy's HP is 30% or below.",
    'Skill5114DescText': 'A 50% chance to increase defense when you recover HP.',
    'Skill5115DescText': 'Restores HP every 12 seconds while in battle.',
    'Skill5117DescText': 'Briefly increases critical rate every 16 seconds while in battle.',
    'Skill5118DescText': 'Briefly increases attack speed every 15 seconds while in battle.',
    'Skill5119DescText': 'Briefly increases defense every 20 seconds while in battle.',
    'Skill5120DescText': 'Briefly increases accuracy every 14 seconds while in battle.',
    'Skill5122DescText': 'Deals heavy bonus damage every 22 seconds while in battle.',
    'Skill5123DescText': 'Increases physical attack.',
    'Skill5124DescText': 'Increases magic attack.',
    'Skill5126DescText': 'Increases defense.',
    'Skill5127DescText': 'Increases critical rate.',
    'Skill5128DescText': 'Increases critical damage.',
    'Skill5129DescText': 'Increases accuracy.',
    'Skill5130DescText': 'Increases attack speed.',
    'Skill5132DescText': 'Increases the amount of HP restored.',
    'Skill5133DescText': 'Increases max memory.',
    'Skill5134DescText': 'Increases physical resistance.',
    'Skill5135DescText': 'Increases experience gained.',
    'Skill5136DescText': 'Increases gold gained.',
    'Skill5137DescText': 'Increases item drop rate.',
    'Skill5138DescText': 'Increases Minimon catch success rate.',
    'Skill5140DescText': 'Increases the number of attracted Minimon you can hold.',
    'Skill5141DescText': 'Reduces potion cooldown.',
    'Skill5143DescText': 'Increases equipment enhancement success rate.',
    'Skill5144DescText': 'Increases movement speed.',
    # ── 능력 이야기 (기존 번역과 같이 둘째 줄은 앞 공백을 유지) ──
    'Skill5012StoryText': 'The depth of a footprint, the grain of a breath.\n A hunter knows by instinct.',
    'Skill5010StoryText': 'When in haste, take the long way around.\n Yet the road keeps growing shorter.',
    'Skill5011StoryText': 'Every road you have walked becomes nourishment.',
    'Skill5013StoryText': 'The mark of one who walked more than one path.',
    'Skill5014StoryText': 'Walking two diverging paths side by side.',
    'Skill5015StoryText': 'The one you meet at the end of every road.',
    # ── 물약 — 원래 한/영 모두 "[아직 내용이 없습니다.]" 자리표시자였다. 2026-09-29 신속물약 1단계 설명의
    #    말투를 이어 한/영 설명을 새로 썼다(한글은 시트에 직접). 여기 값은 그 영문과 같아야 --check 가 통과한다.
    'SPDPotion2DescText': "A stronger brew than the budget one.\nYour steps feel lighter\nand your heart beats a little faster.\n\nStill that suspicious yellow.\nThis time, I didn't ask.",
}

# ─────────────────────────────────────────────────────────────────────────────
# 3) 하드코딩 한글을 대체하는 신규 키 — (키, 한글=기존 리터럴 그대로, 영문)
#    서식({0:0} 등)은 기존 문자열 보간과 한글 출력이 바이트까지 같도록 원문 서식 지정자를 그대로 둔다.
# ─────────────────────────────────────────────────────────────────────────────
NEW_KEYS = [
    # UI_SkillDetailPopup — 발동조건/효과 설명 (BuildTriggerDesc)
    ('AbilityTriggerChanceFormat', '{0:0}% 확률로 {1}', '{0:0}% chance to {1}'),
    ('AbilityTriggerIntervalFormat', '{0:0.#}초마다 {1}', 'Every {0:0.#}s, {1}'),
    ('AbilityCondBattleStartText', '전투 시작 시', 'At battle start'),
    ('AbilityCondOnAttackText', '공격 시', 'On attack'),
    ('AbilityCondAfterAttackText', '공격 후', 'After attacking'),
    ('AbilityCondOnHitText', '타격 시', 'On landing a hit'),
    ('AbilityCondOnCritText', '치명타 시', 'On critical hit'),
    ('AbilityCondOnMissText', '공격이 빗나갈 시', 'When an attack misses'),
    ('AbilityCondSelfEvadeText', '회피 시', 'When you dodge'),
    ('AbilityCondOnDamagedText', '피격 시', 'When you are hit'),
    ('AbilityCondOnHealText', '체력 회복 시', 'When you recover HP'),
    ('AbilityCondLowHpFormat', 'HP {0:0}% 이하 시', 'When HP is {0:0}% or below'),
    # 기존 코드는 아래 둘을 번역 없이 enum 이름("EnemyEvade")으로 찍고 있었다 — 한글 문구를 새로 정한다.
    ('AbilityCondEnemyEvadeText', '상대 회피 시', 'When the enemy dodges'),
    ('AbilityCondEnemyLowHpFormat', '상대 HP {0:0}% 이하 시', 'When enemy HP is {0:0}% or below'),
    ('AbilityContextInBattleText', '전투 중', 'In battle'),
    ('AbilityContextOutOfBattleText', '비전투 중', 'Out of battle'),
    ('AbilityContextAlwaysText', '항시', 'Always'),
    ('AbilityEffectHealFormat', '체력 {0:0} 회복', 'restore {0:0} HP'),
    ('AbilityEffectBonusDamageFormat', '{0:0} 추가 피해', 'deal {0:0} bonus damage'),
    ('AbilityEffectBuffFormat', '{0} +{1:0}', 'gain {0} +{1:0}'),
    ('AbilityEffectDurationFormat', '({0:0.#}초)', '({0:0.#}s)'),
    # 스탯 이름 (영문은 문장 중간에 들어가므로 소문자 — 문장 맨 앞은 코드가 첫 글자만 대문자로)
    ('AbilityStatPhysicalAtkText', '물리공격력', 'physical attack'),
    ('AbilityStatMagicAtkText', '마법공격력', 'magic attack'),
    ('AbilityStatCritRateText', '치명타 확률', 'critical rate'),
    ('AbilityStatCritDamageText', '치명타 피해', 'critical damage'),
    ('AbilityStatAtkSpeedText', '공격속도', 'attack speed'),
    ('AbilityStatAccuracyText', '명중률', 'accuracy'),
    ('AbilityStatDefText', '방어력', 'defense'),
    ('AbilityStatMaxHpText', '최대 체력', 'max HP'),
    ('AbilityStatEvasionText', '회피율', 'evasion'),
    # 다음 등급 미리보기(BuildNextLevelPreview)가 스탯 이름 없이 enum 이름을 찍던 능력들
    ('AbilityStatPhysicalResistText', '물리 저항', 'physical resistance'),
    ('AbilityStatMaxMemoryText', '최대 기억력', 'max memory'),
    ('AbilityStatCatchRateText', '잡기 성공률', 'catch success rate'),
    ('AbilityStatGuaranteedCatchReduceText', '확정 포획 필요 횟수', 'attempts for guaranteed catch'),
    ('AbilityStatExpGainText', '경험치 획득량', 'EXP gain'),
    ('AbilityStatGoldGainText', '골드 획득량', 'gold gain'),
    ('AbilityStatDropRateText', '아이템 드롭률', 'item drop rate'),
    ('AbilityStatInterestMaxText', '관심 미니몬 보관 수', 'attracted Minimon capacity'),
    ('AbilityStatPotionCooldownText', '물약 재사용 대기시간', 'potion cooldown'),
    ('AbilityStatEnhanceLuckText', '강화 성공률', 'enhancement success rate'),
    ('AbilityStatMoveSpeedText', '이동 속도', 'movement speed'),
    ('AbilityStatRecoveryRateText', '회복량', 'HP recovery'),
    # UI_SkillDetailPopup — 기억 버튼 라벨 (다른 버튼 라벨처럼 영문은 대문자)
    ('AbilityNeedLearnText', '습득 필요', 'LEARN FIRST'),
    ('AbilityAlwaysActiveText', '항시 적용', 'ALWAYS ON'),
    # UI_Village1BtnGroup — 대화창 화자 이름
    ('NpcReceptionistNameText', '접수원 아가씨', 'Receptionist'),
    ('NpcAlchemistNameText', '연금공방 아가씨', 'Alchemist'),
    ('NpcBlacksmithNameText', '대장장이', 'Blacksmith'),
    ('NpcMonasteryPriestNameText', '수도원 사제', 'Monastery Priest'),
    ('NpcShopkeeperNameText', '상점 주인', 'Shopkeeper'),
    # 폴백 이름
    ('MinimonDefaultNameText', '미니몬', 'Minimon'),
    ('AbilityUnknownFormat', '능력 #{0}', 'Ability #{0}'),
    # UI_GameScene — 방치 사냥/오프라인 보상 토스트
    ('IdleHuntAreaFormat', '방치 사냥 중 — {0}에서 {1}마리\n경험치 {2:N0} · 골드 {3:N0}',
     'Idle hunting — {1} defeated in {0}\nEXP {2:N0} · Gold {3:N0}'),
    ('IdleHuntFormat', '방치 사냥 중 — {0}마리\n경험치 {1:N0} · 골드 {2:N0}',
     'Idle hunting — {0} defeated\nEXP {1:N0} · Gold {2:N0}'),
    ('OfflineRewardFormat', '자리를 비운 {0} 동안 {1}마리를 처치했습니다\n경험치 {2:N0} · 골드 {3:N0}',
     'You defeated {1} monsters while away for {0}.\nEXP {2:N0} · Gold {3:N0}'),
    ('OfflineRewardCapFormat', '(최대 {0:F0}시간까지 정산됩니다)', '(Rewards are capped at {0:F0} hours.)'),
    ('OfflineTimeHoursFormat', '{0:F1}시간', '{0:F1} hours'),
    ('OfflineTimeMinutesFormat', '{0:F0}분', '{0:F0} min'),
    # GameManager.cs 용 — 코드 교체는 그 파일 담당이 하고, 키/상수만 미리 넣어 둔다.
    #  ("Lv.{0} 필요" 는 기존 RequiredLevelFormat 재사용)
    ('ShopNotForSaleText', '판매하지 않는 물건', 'Not for sale'),
    ('TravelLockedText', '잠김', 'Locked'),
    ('DefaultPlayerNameText', '여행자', 'Traveler'),
]

# 코드가 문자열 리터럴로 Util.Localize("...") 하는 키 — 시트에만 추가하고 Define 상수는 만들지 않는다.
#  영문 표기: 짧은 라벨/버튼은 이웃(START, LANGUAGE, BGM, POPUP TEXT, CONFIRM)처럼 대문자, 문장은 평서문 그대로.
LITERAL_KEYS = [
    ('BackAgainToQuitText', '한 번 더 누르면 게임을 종료합니다.', 'Press back again to quit.'),
    ('QuitGameText', '게임 종료', 'QUIT GAME'),
    ('SettingText', '설정', 'SETTINGS'),
    ('CreditsText', '크레딧', 'CREDITS'),
    ('SystemText', '시스템', 'SYSTEM'),
    ('FullscreenText', '전체 화면', 'FULLSCREEN'),
    ('ThanksForPlayingText', '플레이해 주셔서 감사합니다!', 'Thank you for playing!'),
    ('EquipmentTitleText', '장비', 'EQUIPMENT'),
    # 영어 모드 전수 스캔(2026-09-29)에서 프리팹에 박힌 한국어로 남아 있던 라벨
    ('InventoryTitleText', '소지품', 'INVENTORY'),
    ('InventorySortText', '정리', 'SORT'),
    ('MonsterSkillNeedLinkText', '100% 동기화 필요', 'REQUIRES 100% LINK'),
    ('MonsterSkillFullLockText', '하수인 또는 완벽분석 필요', 'REQUIRES MINIMON OR FULL LOG'),
    ('NotOwnedText', '미보유', 'NOT OWNED'),
    ('InfoBtnText', '정보', 'INFO'),
    ('PotionUpgradeText', '물약개선', 'UPGRADE'),
    ('CustomOrderGreetingText', '"무엇을 만들어드릴까요?"', '"What shall I brew for you?"'),
    # 물약 단계 — 예전엔 '{단계}{RankText}' 를 이어 붙여 영어가 '1RANK' 로 나왔다
    ('RankFormat', '{0}단계', 'RANK {0}'),
    # 물약 상세 전환 버튼(폭 30) — 'REQ. ITEMS' 는 버튼 밖으로 넘친다. 절 제목(ItemNeededText)은 그대로 둔다
    ('SwitchToItemsText', '필요 재료', 'ITEMS'),
    # v2 훈련 패널(ADR0030) — UI_TrainingPanel
    ('TrainingTitleText', '훈련', 'TRAINING'),
    ('TrainingCapFormat', '상한 Lv.{0}', 'CAP Lv.{0}'),
    ('TrainAtkText', '공격력', 'ATK'),
    ('TrainHpText', '체력', 'HP'),
    ('TrainDefText', '방어력', 'DEF'),
    ('TrainCritDmgText', '치명 피해', 'CRIT DMG'),
    ('TrainingRowFormat', '{0}  Lv.{1}  +{2}%', '{0}  Lv.{1}  +{2}%'),
    ('TrainingCostFormat', '{0:N0} 골드', '{0:N0} G'),
    ('TrainingMaxText', '상한', 'MAX'),
    # v2 스테이지·보스(ADR0030) — UI_StageBar / UI_BattleBtnGroup. 알림은 UI_Toast 한 줄에 들어가게 짧게
    ('StageLabelFormat', '{0}-{1}  {2}', '{0}-{1}  {2}'),
    ('StageKillFormat', '처치 {0}/{1}', 'KILLS {0}/{1}'),
    ('StageBossNextText', '다음은 보스!', 'BOSS NEXT!'),
    ('StageBossChallengeText', '눌러서 보스 도전', 'TAP TO FIGHT THE BOSS'),
    ('StageBossTimerFormat', '보스 {0:0.0}초', 'BOSS {0:0.0}s'),
    ('BossAppearText', '보스 출현!', 'BOSS INCOMING!'),
    ('StageClearFormat', '스테이지 돌파! 크리스탈 +{0}', 'STAGE CLEAR! +{0} CRYSTALS'),
    ('BossDefeatText', '보스 처치!', 'BOSS DEFEATED!'),
    ('BossFailText', '보스를 놓쳤다! 훈련 후 재도전', 'BOSS ESCAPED! TRAIN UP'),
    ('StageNextLockedFormat', '다음 스테이지: {0}', 'NEXT STAGE: {0}'),
    # v2 가드·패링(ADR0030) — PlayerController / UI_BattleInput(방패 버튼, ADR0033)
    ('ParryText', '패링!', 'PARRY!'),
    ('BlockText', '막기', 'BLOCK'),
    ('GuardHintText', '\'!\'가 뜨면 방패 버튼으로 막기\n딱 맞추면 패링!', 'SHIELD WHEN "!" SHOWS\nPERFECT TIMING = PARRY!'),
    # v2 탭 공격(ADR0033) — UI_BattleInput
    ('TapHintText', '화면을 누르면 공격! 연타로 몰아치자', 'TAP TO ATTACK! KEEP TAPPING!'),
    ('AutoBattleText', 'AUTO', 'AUTO'),
    ('AutoBattleOnText', '자동 전투 켬 - 누르면 더 빨리 벤다', 'AUTO ON - TAP TO STRIKE FASTER'),
    ('AutoBattleOffText', '자동 전투 끔 - 눌러서 공격', 'AUTO OFF - TAP TO ATTACK'),
    # v2 요일 던전(ADR0034) — UI_MissionPopup 던전 탭 / UI_StageBar / DungeonManager / MissionData
    ('MissionDungeonTabText', '던전', 'DUNGEON'),
    ('DungeonTitleText', '요일 던전', 'DAILY DUNGEONS'),
    ('DungeonGoldText', '골드 던전', 'GOLD DUNGEON'),
    ('DungeonExpText', '경험치 던전', 'EXP DUNGEON'),
    ('DungeonMatText', '재료 던전', 'MATERIAL DUNGEON'),
    ('DungeonTodayFormat', '오늘 {0}배', 'TODAY x{0}'),
    ('DungeonBestFormat', '최고 {0}', 'BEST {0}'),
    ('DungeonFreeFormat', '입장 {0}/{1}', 'ENTER {0}/{1}'),
    ('DungeonHelpFormat', '{0}초 동안 잡은 만큼 보상 · 빨리 누를수록 많이 잡는다', 'REWARDS FOR EVERY KILL IN {0}s - TAP FAST'),
    ('DungeonHudFormat', '남은 {0:0.0}초 · {1}마리', '{0:0.0}s · {1} KILLS'),
    ('DungeonStartFormat', '{0} 시작! {1}초', '{0} - {1}s!'),
    ('DungeonResultFormat', '{0} 끝! {1}마리 · {2}', '{0} DONE! {1} KILLS · {2}'),
    ('DungeonMatRewardFormat', '재료 {0}개 · 만드라고라 {1}', '{0} MATERIALS · {1} MANDRAKE'),
    ('DungeonBusyText', '지금은 들어갈 수 없다', 'CANNOT ENTER RIGHT NOW'),
    ('RewardExpFormat', '경험치 {0:N0}', '{0:N0} EXP'),
    ('MissionDungeonFormat', '요일 던전 {0}회', 'Dungeon runs x{0}'),
    # v2 미션·출석(ADR0030) — UI_MissionPopup / MissionData.TitleKey({0} = 목표)
    ('MissionTitleText', '미션', 'MISSIONS'),
    ('MissionDailyTabText', '일일', 'DAILY'),
    ('MissionWeeklyTabText', '주간', 'WEEKLY'),
    ('MissionGuideTabText', '가이드', 'GUIDE'),
    ('MissionAttendanceTabText', '출석', 'LOGIN'),
    ('MissionKillFormat', '몬스터 {0}마리 처치', 'Defeat monsters x{0}'),
    ('MissionBossFormat', '보스 {0}회 처치', 'Defeat bosses x{0}'),
    ('MissionStageFormat', '스테이지 {0}곳 돌파', 'Clear stages x{0}'),
    ('MissionTrainFormat', '훈련 {0}회', 'Train x{0}'),
    ('MissionCaptureFormat', '포획 {0}회 시도', 'Capture attempts x{0}'),
    ('MissionEnhanceFormat', '장비 강화 {0}회', 'Enhance gear x{0}'),
    ('MissionParryFormat', '패링 {0}회', 'Parry x{0}'),
    ('MissionDailyAllFormat', '일일 미션 완료 {0}일', 'Finish all dailies x{0} days'),
    ('MissionDailyChestText', '일일 미션 모두 완료', 'Finish all daily missions'),
    ('MissionLevelFormat', '레벨 {0} 달성', 'Reach Lv.{0}'),
    ('MissionDiscoverFormat', '도감 {0}종 발견', 'Discover {0} species'),
    ('AttendanceDayFormat', '{0}일째', 'Day {0}'),
    ('MissionClaimText', '받기', 'CLAIM'),
    ('MissionClaimAllText', '모두 받기', 'CLAIM ALL'),
    ('MissionDoneText', '완료', 'DONE'),
    ('MissionNextText', '다음', 'NEXT'),
    ('MissionGuideDoneText', '가이드를 모두 마쳤다!', 'All guide steps done!'),
    ('AttendanceClaimedText', '받음', 'GOT'),
    ('AttendanceTomorrowText', '내일', 'TOMORROW'),
    ('RewardCrystalFormat', '크리스탈 {0}', '{0} CRYSTALS'),
    ('RewardMandrakeFormat', '만드라고라 {0}', '{0} MANDRAKE'),
    ('RewardGoldFormat', '골드 {0:N0}', '{0:N0} GOLD'),
    # v2 상점·광고·핫타임·인앱(ADR0030) — UI_MissionPopup 상점 탭 / ProductData.NameKey
    ('MissionShopTabText', '상점', 'SHOP'),
    ('HotGoldText', '골드 2배', 'GOLD x2'),
    ('HotExpText', '경험치 2배', 'EXP x2'),
    ('HotAtkText', '공격력 +20%', 'ATK +20%'),
    ('HotLeftFormat', '(남은 {0})', '({0} LEFT)'),
    ('HotDurationFormat', '{0}분', '{0} MIN'),
    ('ShopTitleText', '상점', 'SHOP'),
    ('AdWatchFormat', '광고 보기 {0}', 'WATCH AD {0}'),
    ('FreeCrystalFormat', '무료 크리스탈 {0}', 'FREE {0} CRYSTALS'),
    ('OfflineDoubleText', '오프라인 보상 한 번 더', 'DOUBLE OFFLINE REWARD'),
    ('MonthlyDailyFormat', '월정액 오늘 크리스탈 {0} ({1}일 남음)', 'MONTHLY: {0} CRYSTALS TODAY ({1}D LEFT)'),
    ('MonthlyLeftFormat', '{0}일 남음', '{0}D LEFT'),
    ('MandrakeBuyFormat', '만드라고라 1개 (보유 {0})', '1 MANDRAKE (HAVE {0})'),
    ('FirstBuyBonusText', '첫 구매 2배', 'FIRST x2'),
    ('PriceFormat', '{0:N0}원', 'KRW {0:N0}'),
    ('OwnedText', '구매함', 'OWNED'),
    ('StepLockedText', '앞 단계 먼저', 'LOCKED'),
    ('ShopBoughtFormat', '{0} 구매 완료(테스트)', '{0} PURCHASED (TEST)'),
    ('ShopNotReadyText', '결제는 곧 열립니다', 'STORE COMING SOON'),
    ('NotEnoughCrystalText', '크리스탈이 부족합니다', 'NOT ENOUGH CRYSTALS'),
    ('ProductStarterText', '시작 패키지', 'STARTER PACK'),
    ('ProductStepUp1Text', '스텝업 I', 'STEP-UP I'),
    ('ProductStepUp2Text', '스텝업 II', 'STEP-UP II'),
    ('ProductStepUp3Text', '스텝업 III', 'STEP-UP III'),
    ('ProductMonthlyText', '월정액 (30일)', 'MONTHLY PASS (30D)'),
    ('ProductAdFreeText', '광고 제거', 'REMOVE ADS'),
    ('ProductCrystalFormat', '크리스탈 {0}', '{0} CRYSTALS'),
]

# 영문이 Kor 과 같아도 되는 행. 한글 없는 원문(숫자/기호/NEW/MISS)은 자동 허용이라 예외가 생길 때만 키를 적는다.
LEGIT_SAME_KEYS = set()

MASS_ITEMS = {'Spider Silk', 'Wind Essence', 'Damp Moss', 'Revenant Ash'}   # 셀 수 없는 재료 — 복수형 없음


def out(text=''):
    sys.stdout.buffer.write((text + '\n').encode('utf-8', 'replace'))


def short(value, width=260):
    text = '' if value is None else str(value).replace('\n', '\\n')
    return text if len(text) <= width else text[:width - 1] + '…'


# ─────────────────────────────────────────────────────────────────────────────
# 워크북 처리
# ─────────────────────────────────────────────────────────────────────────────
def snapshot(wb):
    return {(ws.title, cell.coordinate): (cell.value, copy(cell._style)) for ws in wb for row in ws for cell in row}


def index_keys(ws):
    rows = {}
    for row in range(2, ws.max_row + 1):
        key = ws.cell(row, 1).value
        if key is None:
            continue
        assert key not in rows, f'{ws.title} 중복 키: {key}'
        rows[key] = row
    return rows


def plural(name):
    if name in MASS_ITEMS:
        return name
    return name[:-1] + 'ies' if name.endswith('y') else name + 's'


def build_name_map(loc_ws):
    """한글명 → (키, 영문명). 같은 한글명이 둘 이상이면 모호하므로 중단."""
    result = {}
    for row in range(2, loc_ws.max_row + 1):
        key, kor, eng = (loc_ws.cell(row, c).value for c in (1, 2, 3))
        if not key or not kor or not eng:
            continue
        name = str(kor).strip()
        if name not in result:
            result[name] = (key, eng)
        elif result[name] is not None and result[name][1] != eng:
            result[name] = None   # 모호 — 실제로 쓰일 때만 replace_names 에서 오류
    return result


def replace_names(text, name_map):
    """영문 대사 속 한글 이름을 영문명으로 치환. 재료명은 개수 앞에서 복수형, 전부 대문자인 이름은 Title Case."""
    def sub(match):
        found = name_map.get(match.group(0))
        assert found, f'LocalizationData 에서 유일하게 매칭되는 영문명이 없음: {match.group(0)!r}'
        key, eng = found
        if eng.isupper():
            eng = eng.title()   # 'WHISPERING CAVE' → 'Whispering Cave' (문장 중간이라 대문자 표기는 어색)
        if key.startswith('Mat'):
            count = re.search(r'<b>(\d+)</b></color> <color=#2DFF00><b>$', text[:match.start()])
            assert count, f'재료명 앞에서 개수를 찾지 못함: {text!r}'
            if int(count.group(1)) != 1:
                eng = plural(eng)
        return eng

    return HANGUL_RUN.sub(sub, text)


def placeholders(text):
    return sorted(set(re.findall(r'\{(\d+)[^}]*\}', text)))


def plan(loc_wb, dlg_wb):
    """메모리의 워크북에만 변경을 적용하고 [(파일, 시트, 셀, 전, 후)] 를 돌려준다. 재호출하면 변경 0."""
    changes = []

    def patch(file, cell, value):
        if cell.value != value:
            changes.append((file, cell.parent.title, cell.coordinate, cell.value, value))
            cell.value = value

    ws = loc_wb.active
    rows = index_keys(ws)
    last_row = ws.max_row

    for key, eng in SKILL_ENG.items():
        assert key in rows, f'LocalizationData 에 키가 없음: {key}'
        cell = ws.cell(rows[key], 3)
        assert cell.value == eng or HANGUL.search(str(cell.value)), f'{key}: 영문이 이미 다른 번역: {cell.value!r}'
        patch('LocalizationData', cell, eng)

    for key, kor, eng in NEW_KEYS + LITERAL_KEYS:
        assert placeholders(kor) == placeholders(eng), f'{key}: 한/영 자리표시자 불일치'
        assert not HANGUL.search(eng), f'{key}: 영문에 한글'
        if key in rows:
            assert (ws.cell(rows[key], 2).value, ws.cell(rows[key], 3).value) == (kor, eng), f'{key}: 기존 값과 다름'
            continue
        row = ws.max_row + 1
        for column, value in enumerate((key, kor, eng), 1):
            cell = ws.cell(row, column, value)
            cell._style = copy(ws.cell(last_row, column)._style)
            changes.append(('LocalizationData', ws.title, cell.coordinate, None, value))
        rows[key] = row

    name_map = build_name_map(ws)
    for dws in dlg_wb.worksheets:
        for row in range(2, dws.max_row + 1):
            cell = dws.cell(row, 3)
            if isinstance(cell.value, str) and HANGUL.search(cell.value):
                fixed = replace_names(cell.value, name_map)
                assert not HANGUL.search(fixed), f'{dws.title}!{cell.coordinate}: 치환 후에도 한글이 남음'
                patch('DialogueData', cell, fixed)
    return changes


def verify_only_intended(before, wb, changes, file):
    """계획한 셀 말고는 값·서식이 하나도 바뀌지 않았는지 전 시트를 셀 단위로 대조."""
    touched = {(sheet, coord) for f, sheet, coord, _, _ in changes if f == file}
    after = snapshot(wb)
    for pos, (value, style) in before.items():
        if pos in touched:
            continue
        cell = wb[pos[0]][pos[1]]
        assert cell.value == value and cell._style == style, f'{file} {pos}: 계획 밖 변경'
    assert set(after) - set(before) <= touched, f'{file}: 계획 밖 셀 추가'


def sync_define(write):
    """Define.LocalizeKey 에 없는 신규 키 상수를 마지막 상수 뒤에 붙인다(BOM/줄바꿈 방식 유지)."""
    text = DEFINE.read_bytes().decode('utf-8')   # BOM 은 ﻿ 로 그대로 남고 되쓸 때 복원된다
    start = text.index('    #region LocalizeKey')
    end = text.index('    #endregion', start)
    region = text[start:end]
    existing = set(re.findall(r'public const string (\w+) = ', region))
    missing = [key for key, _, _ in NEW_KEYS if key not in existing]
    if missing and write:
        last = list(re.finditer(r'public const string \w+ = "[^"]*";(\r?\n)', region))[-1]
        newline = last.group(1)
        block = ''.join(f'        public const string {key} = "{key}";{newline}' for key in missing)
        at = start + last.end()
        DEFINE.write_bytes((text[:at] + block + text[at:]).encode('utf-8'))
    return missing


# ─────────────────────────────────────────────────────────────────────────────
# --check
# ─────────────────────────────────────────────────────────────────────────────
def strip_cs(src):
    """주석/문자열/문자 리터럴을 지운 코드만 남긴다(중괄호 균형 검사용)."""
    out_chars, i, n = [], 0, len(src)
    while i < n:
        c, nxt = src[i], src[i + 1] if i + 1 < n else ''
        if c == '/' and nxt == '/':
            while i < n and src[i] != '\n':
                i += 1
        elif c == '/' and nxt == '*':
            j = src.find('*/', i + 2)
            i = n if j < 0 else j + 2
        elif c in '$@' and (nxt == '"' or (nxt in '$@' and src[i + 2:i + 3] == '"')):
            verbatim = '@' in src[i:i + 2]
            i += 2 if nxt == '"' else 3
            while i < n:
                if verbatim and src[i] == '"' and src[i + 1:i + 2] == '"':
                    i += 2
                elif not verbatim and src[i] == '\\':
                    i += 2
                elif src[i] == '"':
                    i += 1
                    break
                else:
                    i += 1
        elif c == '"':
            i += 1
            while i < n and src[i] != '"':
                i += 2 if src[i] == '\\' else 1
            i += 1
        elif c == "'":
            i += 1
            while i < n and src[i] != "'":
                i += 2 if src[i] == '\\' else 1
            i += 1
        else:
            out_chars.append(c)
            i += 1
    return ''.join(out_chars)


def check():
    problems = []
    loc = openpyxl.load_workbook(LOC_PATH, read_only=True)
    dlg = openpyxl.load_workbook(DLG_PATH, read_only=True)
    try:
        rows = list(loc.active.iter_rows(min_row=2, values_only=True))
        keys = [r[0] for r in rows if r[0] is not None]
        dup = {k for k in keys if keys.count(k) > 1}
        hangul_eng = [r[0] for r in rows if r[0] and r[2] and HANGUL.search(str(r[2]))]
        empty_eng = [r[0] for r in rows if r[0] and not (r[2] and str(r[2]).strip())]
        same = [r[0] for r in rows if r[0] and r[1] and r[2] and HANGUL.search(str(r[1]))
                and str(r[1]).strip() == str(r[2]).strip() and r[0] not in LEGIT_SAME_KEYS]
        out(f'Localization: 행 {len(keys)}, 중복 키 {len(dup)}, 영문 빈 행 {len(empty_eng)}, '
            f'영문에 한글 {len(hangul_eng)}, 영문==한글(한글 있는 원문) {len(same)}')
        for label, bad in (('중복 키', dup), ('영문 빈 행', empty_eng), ('영문에 한글', hangul_eng), ('영문==한글', same)):
            if bad:
                problems.append(f'Localization {label}: {sorted(bad)[:10]}')

        lines = dlg_hangul = dlg_same = 0
        for ws in dlg.worksheets:
            for key, kor, eng in ws.iter_rows(min_row=2, max_col=3, values_only=True):
                if key is None:
                    continue
                lines += 1
                dlg_hangul += bool(eng and HANGUL.search(str(eng)))
                dlg_same += bool(eng and kor and HANGUL.search(str(kor)) and str(eng).strip() == str(kor).strip())
        out(f'Dialogue: 대사 {lines}줄, 영문에 한글 {dlg_hangul}, 영문==한글 {dlg_same}')
        if dlg_hangul or dlg_same:
            problems.append('Dialogue 에 미번역 영문 대사가 남음')

        sheet_rows = {r[0]: r for r in rows if r[0]}
        for key, kor, eng in NEW_KEYS + LITERAL_KEYS:
            row = sheet_rows.get(key)
            if not row or (row[1], row[2]) != (kor, eng):
                problems.append(f'신규 키가 시트에 없거나 값이 다름: {key}')
        out(f'신규 키 시트 반영: 상수 대상 {len(NEW_KEYS)}개 + 문자열 리터럴 대상 {len(LITERAL_KEYS)}개')
        for key, eng in SKILL_ENG.items():
            row = sheet_rows.get(key)
            if not row or row[2] != eng:
                problems.append(f'영문 채움이 반영되지 않음: {key}')
    finally:
        loc.close()
        dlg.close()

    # Define.LocalizeKey 상수 ↔ 시트
    define = DEFINE.read_bytes().decode('utf-8')
    region = define[define.index('    #region LocalizeKey'):define.index('    #endregion', define.index('    #region LocalizeKey'))]
    consts = dict(re.findall(r'public const string (\w+) = "([^"]*)";', region))
    added_missing = [key for key, _, _ in NEW_KEYS if consts.get(key) != key or key not in sheet_rows]
    stale = [name for name, value in consts.items() if value not in sheet_rows]
    out(f'Define.LocalizeKey: 상수 {len(consts)}개, 신규 {len(NEW_KEYS)}개 중 누락/불일치 {len(added_missing)}, '
        f'시트에 없는 기존 상수 {len(stale)}')
    if added_missing:
        problems.append(f'신규 키 상수 누락/불일치: {added_missing}')
    if stale:
        out(f'  (참고: 시트에 없는 상수 — 이 작업과 무관하면 무시) {stale[:8]}')

    # C# 중괄호 균형 + LocalizeKey 참조 정의 여부
    for rel in CS_FILES:
        src = (SCRIPTS / rel).read_bytes().decode('utf-8-sig')
        code = strip_cs(src)
        refs = sorted(set(re.findall(r'LocalizeKey\.(\w+)', code)))
        undefined = [r for r in refs if r not in consts]
        dangling = [r for r in refs if r in consts and consts[r] not in sheet_rows]
        opens, closes = code.count('{'), code.count('}')
        out(f'{rel}: {{ {opens} / }} {closes} {"OK" if opens == closes else "불균형"}, '
            f'LocalizeKey 참조 {len(refs)}개, 미정의 {len(undefined)}, 시트에 없음 {len(dangling)}')
        if opens != closes:
            problems.append(f'{rel} 중괄호 불균형')
        if undefined or dangling:
            problems.append(f'{rel} LocalizeKey 문제: 미정의 {undefined}, 시트에 없음 {dangling}')

    out('CHECK ' + ('FAIL' if problems else 'PASS'))
    for problem in problems:
        out(f'  - {problem}')
    return not problems


# ─────────────────────────────────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument('--write', action='store_true', help='변경을 저장 (기본은 미리보기)')
    mode.add_argument('--check', action='store_true', help='저장된 결과 재검증')
    args = parser.parse_args()
    if args.check:
        sys.exit(0 if check() else 1)

    loc_wb = openpyxl.load_workbook(LOC_PATH)
    dlg_wb = openpyxl.load_workbook(DLG_PATH)
    try:
        before = {'LocalizationData': snapshot(loc_wb), 'DialogueData': snapshot(dlg_wb)}
        changes = plan(loc_wb, dlg_wb)
        verify_only_intended(before['LocalizationData'], loc_wb, changes, 'LocalizationData')
        verify_only_intended(before['DialogueData'], dlg_wb, changes, 'DialogueData')
        assert not plan(loc_wb, dlg_wb), '같은 데이터 재적용에서 변경이 또 발생(멱등성 위반)'

        counts = {}
        for file, sheet, coord, old, new in changes:
            kind = ('추가' if old is None else '수정')
            counts[(file, kind)] = counts.get((file, kind), 0) + 1
            out(f'{file} {sheet}!{coord} [{kind}] {short(old)!r} -> {short(new)!r}')
        added_consts = sync_define(args.write)
        for (file, kind), n in sorted(counts.items()):
            out(f'{file}: {kind} 셀 {n}개')
        out(f'Define.cs: 추가할 상수 {len(added_consts)}개')
        if not args.write:
            out('미리보기입니다 — 저장하려면 --write')
            return
        if any(f == 'LocalizationData' for f, *_ in changes):
            loc_wb.save(LOC_PATH)
        if any(f == 'DialogueData' for f, *_ in changes):
            dlg_wb.save(DLG_PATH)
        out('저장 완료 — Unity Tools/ParseAllLocal 로 JSON 을 재생성하세요.')
    finally:
        loc_wb.close()
        dlg_wb.close()


if __name__ == '__main__':
    main()
