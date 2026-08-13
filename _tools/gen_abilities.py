# -*- coding: utf-8 -*-
"""능력(스킬) 콘텐츠 확충 — 기존 20종에 45종 추가.

기획서 근거
  v3 슬38: 능력 종류(유틸/전투/편의), 기억력 점유량, 재능(점유량 0),
           획득 5경로(장비+9 / 미니몬 공명도100% / NPC재료 / 전직 / 도감별)
  v3 슬39: 전투 발동 조건 — 확률 발동 15종 + 시간 발동(쿨타임)

설계
  - 확률 발동 15조건을 **전부 한 번씩** 쓰는 전투 능력 15종 (조건이 코드에만 있고
    쓰는 능력이 없으면 시스템이 죽은 채로 남는다)
  - 시간 발동(쿨타임) 8종
  - 스탯 패시브 12종 — 스탯 종류별로 하나씩, 기억력 점유량으로 취사선택하게
  - 유틸/편의 10종 — 경험치·골드·포획·관심 등 전투 밖 성장축
  획득 경로는 5종에 고르게 배분해 어느 한 경로만 파도 다 얻지 못하게 한다.

기억력 점유량: 효과가 클수록 크게(1~5). 재능(0)은 항시 적용이라 소수만.
실행: python _tools/gen_abilities.py → 이후 Unity에서 Tools/ParseAllLocal
"""
import os
import openpyxl

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
EXCEL = os.path.join(BASE, "Assets", "@Resources", "Data", "Excel")
START_ID = 5100

# (조건, 한글 능력명, 설명, 효과, 효과스탯, 효과값, 지속, 확률, 조건값, 점유, 출처, 최대레벨)
COMBAT = [
    ("BattleStart",      "선제의 감각", "전투 시작 시 {p}% 확률로 공격력이 잠시 오른다",       "Buff", "PhysicalAtk", 25, 8,  60, 0,  2, "Class",     3),
    ("OnAttack",         "연격의 습관", "공격 시 {p}% 확률로 추가 피해를 준다",                 "BonusDamage", "", 30, 0, 25, 0,  2, "Equip",     5),
    ("AfterAttack",      "잔심",       "공격 후 {p}% 확률로 다음 일격이 날카로워진다",         "Buff", "CritRate",   12, 5,  30, 0,  2, "Npc",       3),
    ("OnHit",            "타격 감각",   "타격 시 {p}% 확률로 체력을 조금 회복한다",             "Heal", "",           18, 0,  20, 0,  2, "Minion",    3),
    ("OnCrit",           "치명의 흐름", "치명타 시 {p}% 확률로 공격속도가 오른다",              "Buff", "AtkSpeed",   20, 6,  45, 0,  2, "Unlimited", 4),
    ("ConsecutiveCrit",  "연쇄 각성",   "연속 {c}회 치명타 시 큰 피해를 추가한다",              "BonusDamage", "", 90, 0, 100, 3,  3, "Class",     3),
    ("OnMiss",           "빗나감의 교훈", "빗나갈 때 {p}% 확률로 명중이 오른다",                "Buff", "Accuracy",   15, 8,  70, 0,  1, "Npc",       3),
    ("EnemyEvade",       "간파",       "상대가 회피하면 {p}% 확률로 치명타율이 오른다",        "Buff", "CritRate",   15, 6,  50, 0,  2, "Equip",     3),
    ("SelfEvade",        "흘리기",     "회피에 성공하면 {p}% 확률로 반격 피해를 준다",         "BonusDamage", "", 45, 0, 55, 0,  2, "Minion",    4),
    ("ConsecutiveEvade", "그림자 걸음", "연속 {c}회 회피 시 회피율이 크게 오른다",              "Buff", "Evasion",    18, 10, 100, 3, 3, "Class",     3),
    ("OnDamaged",        "굳은 살",     "피격 시 {p}% 확률로 방어가 단단해진다",                "Buff", "Def",        25, 6,  35, 0,  2, "Equip",     4),
    ("LowHp",            "벼랑 끝",     "체력이 {c}% 이하가 되면 공격력이 크게 오른다",         "Buff", "PhysicalAtk", 60, 12, 100, 30, 4, "Unlimited", 3),
    ("EnemyLowHp",       "숨통",       "상대 체력이 {c}% 이하일 때 추가 피해를 준다",          "BonusDamage", "", 70, 0, 100, 30, 3, "Npc",       4),
    ("FatalDamage",      "불굴",       "치명상을 입을 때 {p}% 확률로 체력을 회복하며 버틴다",  "Heal", "",          200, 0,  40, 0,  5, "Class",     2),
    ("OnHeal",           "회복의 여운", "회복할 때 {p}% 확률로 방어가 오른다",                  "Buff", "Def",        18, 8,  50, 0,  1, "Minion",    3),
]

# 시간 발동(쿨타임) — (한글명, 설명, 효과, 스탯, 값, 지속, 쿨다운, 점유, 출처, 최대레벨)
TIMED = [
    ("전투 호흡",   "{cd}초마다 체력을 회복한다",             "Heal", "",            22, 0,  12, 2, "Npc",       4),
    ("각성",       "{cd}초마다 공격력이 잠시 오른다",         "Buff", "PhysicalAtk", 30, 6,  18, 3, "Class",     4),
    ("집중",       "{cd}초마다 치명타율이 잠시 오른다",       "Buff", "CritRate",    18, 6,  16, 2, "Equip",     4),
    ("가속",       "{cd}초마다 공격속도가 잠시 오른다",       "Buff", "AtkSpeed",    25, 5,  15, 2, "Minion",    4),
    ("철벽",       "{cd}초마다 방어가 잠시 오른다",           "Buff", "Def",         35, 6,  20, 2, "Equip",     3),
    ("예리함",     "{cd}초마다 명중이 잠시 오른다",           "Buff", "Accuracy",    20, 6,  14, 1, "Unlimited", 3),
    ("잔상",       "{cd}초마다 회피율이 잠시 오른다",         "Buff", "Evasion",     16, 6,  17, 2, "Class",     3),
    ("파쇄",       "{cd}초마다 강한 추가 피해를 준다",        "BonusDamage", "",     80, 0,  22, 3, "Npc",       5),
]

# 스탯 패시브 — (스탯, 한글명, 설명, 보너스, 점유, 출처, 최대레벨)
PASSIVE = [
    ("PhysicalAtk", "완력",       "물리 공격력이 오른다",       12, 2, "Equip",     5),
    ("MagicAtk",    "마력 회로",   "마법 공격력이 오른다",       10, 2, "Npc",       5),
    ("MaxHp",       "강인함",     "최대 체력이 오른다",        120, 2, "Class",     5),
    ("Def",         "단련",       "방어력이 오른다",            10, 2, "Equip",     5),
    ("CritRate",    "급소 감각",   "치명타 확률이 오른다",        4, 3, "Minion",    5),
    ("CritDamage",  "일격필살",    "치명타 피해가 오른다",       15, 3, "Unlimited", 5),
    ("Accuracy",    "정확",       "명중률이 오른다",             5, 1, "Npc",       5),
    ("AtkSpeed",    "속공",       "공격속도가 오른다",           6, 3, "Minion",    5),
    ("Evasion",     "민첩",       "회피율이 오른다",             4, 2, "Class",     5),
    ("RecoveryRate", "치유 친화",  "회복량이 오른다",            10, 1, "Npc",       4),
    ("MaxMemory",   "기억의 그릇", "최대 기억력이 오른다",        3, 0, "Unlimited", 5),
    ("PhysicalResist", "감내",     "물리 저항이 오른다",          5, 2, "Equip",     4),
]

# 유틸/편의 — (스탯키, 한글명, 설명, 값, 점유, 출처, 카테고리, 최대레벨)
UTILITY = [
    ("ExpGain",              "탐구심",     "경험치 획득량이 늘어난다",           8, 2, "Unlimited", "Utility",     5),
    ("GoldGain",             "장사 수완",   "골드 획득량이 늘어난다",            10, 1, "Npc",       "Utility",     5),
    ("DropRate",             "행운",       "아이템 드롭률이 오른다",             6, 2, "Minion",    "Utility",     5),
    ("CatchRate",            "교감",       "미니몬 잡기 성공률이 오른다",        5, 2, "Minion",    "Utility",     5),
    ("GuaranteedCatchReduce", "각인",      "확정 포획에 필요한 횟수가 줄어든다",  1, 3, "Class",     "Utility",     3),
    ("InterestMax",          "친화의 향기", "관심 미니몬 보관 수가 늘어난다",      1, 2, "Npc",       "Convenience", 3),
    ("PotionCooldown",       "약초 지식",   "물약 재사용 대기시간이 줄어든다",     8, 2, "Npc",       "Convenience", 4),
    ("ComboKeep",            "연계 감각",   "콤보 게이지가 더 오래 유지된다",     10, 2, "Class",     "Convenience", 4),
    ("EnhanceLuck",          "대장장이의 눈", "장비 강화 성공률이 오른다",         4, 3, "Equip",     "Utility",     3),
    ("MoveSpeed",            "발걸음",     "이동 속도가 오른다",                 6, 1, "Unlimited", "Convenience", 3),
]


def rows():
    out = []
    did = START_ID
    for cond, kor, desc, eff, estat, eval_, edur, proc, cval, mem, src, maxlv in COMBAT:
        out.append(dict(id=did, kor=kor, desc=desc.format(p=proc, c=cval), cat="Combat",
                        active="Passive", stat="", bonus=0, maxlv=maxlv, mem=mem, src=src,
                        trigger="Probability", conds=cond, proc=proc, cd=0, cval=cval,
                        eff=eff, estat=estat, eval_=eval_, edur=edur))
        did += 1
    for kor, desc, eff, estat, eval_, edur, cd, mem, src, maxlv in TIMED:
        out.append(dict(id=did, kor=kor, desc=desc.format(cd=cd), cat="Combat",
                        active="Passive", stat="", bonus=0, maxlv=maxlv, mem=mem, src=src,
                        trigger="Time", conds="", proc=100, cd=cd, cval=0,
                        eff=eff, estat=estat, eval_=eval_, edur=edur))
        did += 1
    for stat, kor, desc, bonus, mem, src, maxlv in PASSIVE:
        out.append(dict(id=did, kor=kor, desc=desc, cat="Memory", active="Passive",
                        stat=stat, bonus=bonus, maxlv=maxlv, mem=mem, src=src,
                        trigger="None", conds="", proc=100, cd=0, cval=0,
                        eff="None", estat="", eval_=0, edur=0))
        did += 1
    for stat, kor, desc, val, mem, src, cat, maxlv in UTILITY:
        out.append(dict(id=did, kor=kor, desc=desc, cat=cat, active="Passive",
                        stat=stat, bonus=val, maxlv=maxlv, mem=mem, src=src,
                        trigger="None", conds="", proc=100, cd=0, cval=0,
                        eff="None", estat="", eval_=0, edur=0))
        did += 1
    return out


def write_skill(rs):
    p = os.path.join(EXCEL, "SkillData.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb[wb.sheetnames[0]]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    have = {ws.cell(r, h["DataId"]).value for r in range(4, ws.max_row + 1)}
    row = ws.max_row + 1
    added = 0
    for d in rs:
        if d["id"] in have:
            continue
        vals = {
            "DataId": d["id"], "Name": f'Skill{d["id"]}NameText', "Description": f'Skill{d["id"]}DescText',
            "IconLabel": f'Skill{d["id"]}', "Category": d["cat"], "SkillActiveType": d["active"],
            "StatType": d["stat"], "StatBonus": d["bonus"], "MaxLevel": d["maxlv"],
            "MemoryCapacity": d["mem"], "Source": d["src"], "RequiredMinionId": 0,
            "RequiredEquipId": 0, "RequiredItemId": 0, "RequiredItemCount": 0,
            "TriggerType": d["trigger"], "Conditions": d["conds"], "ProcChance": d["proc"],
            "Cooldown": d["cd"], "ActiveContext": "InBattle", "ConditionValue": d["cval"],
            "Effect": d["eff"], "EffectStat": d["estat"], "EffectValue": d["eval_"],
            "EffectDuration": d["edur"],
        }
        for name, col in h.items():
            if name in vals:
                ws.cell(row, col, vals[name])
        row += 1
        added += 1
    wb.save(p)
    print(f"SkillData.xlsx: +{added}종 (총 {len(have) + added}종)")


def write_localization(rs):
    p = os.path.join(EXCEL, "LocalizationData.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb[wb.sheetnames[0]]
    have = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
    row = ws.max_row + 1
    n = 0
    for d in rs:
        for key, kor in ((f'Skill{d["id"]}NameText', d["kor"]), (f'Skill{d["id"]}DescText', d["desc"])):
            if key in have:
                continue
            ws.cell(row, 1, key); ws.cell(row, 2, kor); ws.cell(row, 3, kor)
            have.add(key); row += 1; n += 1
    wb.save(p)
    print(f"LocalizationData.xlsx: 능력 텍스트 +{n}개")


def main():
    rs = rows()
    write_skill(rs)
    write_localization(rs)
    print(f"\n구성: 전투(확률) {len(COMBAT)} / 전투(쿨타임) {len(TIMED)} / "
          f"스탯 패시브 {len(PASSIVE)} / 유틸·편의 {len(UTILITY)} = {len(rs)}종")
    srcs = {}
    for d in rs:
        srcs[d["src"]] = srcs.get(d["src"], 0) + 1
    print("획득 경로 분포:", srcs)


if __name__ == "__main__":
    main()
