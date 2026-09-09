"""포획 → 동행 능력 체험 → 공명도 100% 영구 습득을 연결한다.

기존 능력/번역/포획 확률은 보존한다. 엑셀만 수정하며 이후 Tools/ParseAllLocal 필수.
--apply 없이 실행하면 변경 예정 개수만 출력한다. ADR0014 참조.
"""
import argparse
from copy import copy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
EXCEL = ROOT / 'Assets/@Resources/Data/Excel'

# 첫 포획까지의 대기만 줄이고, 재포획의 확률과 4회 공명 목표는 그대로 둔다.
FIRST_INTEREST_KILLS = 3
FIRST_CATCH_CHANCE = 1.0
SKILL_OFFSET = 4000

# (효과, 효과 기준/스탯, 값, 쿨타임, 지속시간, 역할, 한국어 효과, 영어 효과)
# 공격·회복은 플레이어 능력치의 %여서 지역을 넘어도 가치가 유지된다.
PROFILES = {
    'heal': ('Heal', 'MaxHp', 18, 18, 0, 'Def',
             '최대 체력의 18%를 회복한다.', 'Restore 18% of maximum HP.'),
    'strike': ('BonusDamage', 'PhysicalAtk', 180, 12, 0, 'Skill',
               '물리 공격력의 180%만큼 확정 추가 피해를 준다.', 'Deal guaranteed bonus damage equal to 180% of physical attack.'),
    'heavy': ('BonusDamage', 'PhysicalAtk', 300, 20, 0, 'Skill',
              '물리 공격력의 300%만큼 확정 추가 피해를 준다.', 'Deal guaranteed bonus damage equal to 300% of physical attack.'),
    'haste': ('Buff', 'AtkSpeed', 30, 18, 8, 'Skill',
              '8초 동안 공격속도가 30%p 오른다.', 'Gain 30 percentage points of attack speed for 8 seconds.'),
    'guard': ('Buff', 'PhysicalResist', 25, 18, 8, 'Def',
              '8초 동안 물리 저항이 25%p 오른다.', 'Gain 25 percentage points of physical resistance for 8 seconds.'),
    'evade': ('Buff', 'Evasion', 18, 18, 8, 'Def',
              '8초 동안 회피율이 18%p 오른다.', 'Gain 18 percentage points of evasion for 8 seconds.'),
    'focus': ('Buff', 'Accuracy', 20, 14, 8, 'Skill',
              '8초 동안 명중률이 20%p 오른다.', 'Gain 20 percentage points of accuracy for 8 seconds.'),
}
PROFILE_BY_BASE = {
    '슬라임': 'heal', '숲거미': 'focus', '들늑대': 'haste', '독버섯': 'strike',
    '바람정령': 'strike', '멧돼지': 'heavy', '까마귀': 'focus', '덩굴손': 'guard',
    '여우': 'evade', '말벌': 'haste', '동굴박쥐': 'focus', '두더지': 'strike',
    '광부망령': 'heavy', '수정골렘': 'guard', '지네': 'haste', '이끼괴물': 'heal',
    '석상병': 'guard', '그림자': 'evade', '도롱뇽': 'strike', '광충': 'focus',
    '망령': 'strike', '석상기사': 'guard', '타락사제': 'heal', '키메라': 'heavy',
    '가고일': 'heavy', '수호자': 'guard', '환영': 'evade', '타락정령': 'strike',
    '제단수호물': 'heavy', '공허체': 'heavy', '고블린': 'strike',
}
UI_TEXT = {
    'CompanionAbilityBorrowedText': ('동행 중 사용', 'In use'),
    'CompanionAbilityAvailableText': ('동행 시 사용', 'Summon to use'),
}


def records(ws, start=4):
    headers = [c.value for c in ws[1]]
    return {ws.cell(r, 1).value: (r, dict(zip(headers, [c.value for c in ws[r]])))
            for r in range(start, ws.max_row + 1) if ws.cell(r, 1).value is not None}


def append_record(ws, values):
    row = ws.max_row + 1
    for col, head in enumerate(ws[1], 1):
        cell = ws.cell(row, col, values.get(head.value))
        if row > 4:
            cell._style = copy(ws.cell(4, col)._style)


def generate(apply=False):
    books = {n: openpyxl.load_workbook(EXCEL / (n + '.xlsx'))
             for n in ('CreatureData', 'MinMonData', 'SkillData', 'LocalizationData')}
    creature = records(books['CreatureData'].active)
    skills_ws = books['SkillData'].active
    skills = records(skills_ws)
    loc_ws = books['LocalizationData'].active
    loc = records(loc_ws, 2)
    min_ws = books['MinMonData'].active
    minions = records(min_ws)
    first_habitat = creature[2001][1]['Habitat']
    translations = dict(UI_TEXT)
    generated = 0
    for mid, (_, m) in minions.items():
        sid = mid + SKILL_OFFSET
        if sid in skills:
            continue
        c = creature[m['MonsterDataId']][1]
        name = loc[c['Name']][1]['Kor']
        base = next((b for b in sorted(PROFILE_BY_BASE, key=len, reverse=True) if name.endswith(b)), None)
        if base is None:
            raise ValueError(f'능력 역할이 없는 마물: {mid} {name}')
        effect, stat, value, cooldown, duration, role, kor, eng = PROFILES[PROFILE_BY_BASE[base]]
        desc_key = f'MinionAbility{mid}DescText'
        translations[desc_key] = (
            f'{kor}\n재사용 {cooldown}초(미니몬 능력끼리 공유). 단계마다 효과량 증가.\n동행 중 사용 가능. 공명도 100%에서 영구 습득.',
            f'{eng}\nCooldown: {cooldown}s, shared by minion abilities. Effect scales with ability level.\nAvailable while accompanying. Permanently learned at 100% resonance.')
        append_record(skills_ws, {
            'DataId': sid, 'Name': c['UniqueSkillName'], 'Description': desc_key,
            'Story': c['UniqueSkillDesc'], 'IconLabel': c['IconLabel'], 'Category': 'Combat',
            'SkillActiveType': 'Repeat', 'StatType': '', 'StatBonus': 0,
            'MaxLevel': 2, 'MemoryCapacity': 3, 'Source': 'Minion',
            'RequiredMinionId': mid, 'RequiredEquipId': 0, 'RequiredItemId': 0,
            'RequiredItemCount': 0, 'RoleTypes': role, 'TriggerType': 'Active',
            'ProcChance': 100, 'Cooldown': cooldown, 'ActiveContext': 'InBattle',
            'ConditionValue': 0, 'Effect': effect, 'EffectStat': stat, 'EffectValue': value,
            'EffectDuration': duration,
            'EffectVfxLabel': 'HP Potion Ef' if effect == 'Heal' else 'FX_hit_10' if effect == 'BonusDamage' else 'Star_01_Anim_13_a',
        })
        generated += 1
    for key, (kor, eng) in translations.items():
        if key not in loc:
            loc_ws.append([key, kor, eng])
        elif key in UI_TEXT:
            # 이 생성기가 소유한 새 UI 문구만 갱신한다. 기존 번역은 보존한다.
            row = loc[key][0]
            loc_ws.cell(row, 2, kor)
            loc_ws.cell(row, 3, eng)
    # 기존 컬럼 뒤에만 추가: 인덱스 기반 Unity 파서와 기존 14열을 보존한다.
    for col, header, kind, desc in ((15, 'FirstInterestKillCount', 'int', '미보유 첫 관심 보장 처치수(0=미사용)'),
                                     (16, 'FirstCatchChance', 'float', '첫 포획 확률(0=기존 확률)')):
        existing = min_ws.cell(1, col).value
        assert existing in (None, header), (col, existing)
        for r, v in enumerate((header, kind, desc), 1):
            min_ws.cell(r, col, v)
    starters = []
    for mid, (row, m) in minions.items():
        early = creature[m['MonsterDataId']][1]['Habitat'] == first_habitat
        min_ws.cell(row, 15, FIRST_INTEREST_KILLS if early else 0)
        min_ws.cell(row, 16, FIRST_CATCH_CHANCE if early else 0)
        if early:
            starters.append(mid)
    print(f'고유능력 추가 {generated}, 첫 지역 {starters}, apply={apply}')
    if apply:
        for name in ('SkillData', 'MinMonData', 'LocalizationData'):
            books[name].save(EXCEL / (name + '.xlsx'))
    for wb in books.values():
        wb.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--apply', action='store_true')
    generate(parser.parse_args().apply)
