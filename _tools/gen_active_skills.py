# -*- coding: utf-8 -*-
"""주인공 액티브 스킬 6종 추가 — 스킬바(퀵슬롯 6칸)에 올라갈 데이터.

기존 능력 65종은 전부 Passive(자동 발동)라 플레이어가 **직접 쓸 수 있는 게 없었다**.
AbilityRunner의 효과 처리(Heal/BonusDamage/Buff)를 그대로 재사용하고,
발동 주체만 플레이어로 바꾼 것이 TriggerType=Active다.

쿨타임은 짧은 딜링기(8초) ~ 필살기(45초)로 벌려 슬롯 6칸이 서로 다른 리듬을 갖게 했다.
"""
import os

import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XL = os.path.join(BASE, "Assets", "@Resources", "Data", "Excel")

# (id, 한글명, 설명, 효과, 스탯, 수치, 지속, 쿨, VFX)
SKILLS = [
    (5200, "강타", "적에게 즉시 추가 피해를 준다.",
     "BonusDamage", "", 40, 0, 8, "FX_hit_10"),
    (5201, "회복술", "체력을 회복한다.",
     "Heal", "", 60, 0, 20, "HP Potion Ef"),
    (5202, "광폭화", "일정 시간 공격력이 크게 오른다.",
     "Buff", "PhysicalAtk", 15, 10, 30, "Star_01_Anim_13_a"),
    (5203, "질풍", "일정 시간 공격 속도가 오른다.",
     "Buff", "AtkSpeed", 20, 8, 25, "Star_01_Anim_13_a"),
    (5204, "철벽", "일정 시간 방어력이 크게 오른다.",
     "Buff", "Def", 12, 10, 25, "Star_01_Anim_13_a"),
    (5205, "필살", "온 힘을 모아 결정타를 날린다.",
     "BonusDamage", "", 150, 0, 45, "FX_hit_10"),
]

ENG = {5200: ("Smash", "Deals instant bonus damage."),
       5201: ("Mend", "Restores health."),
       5202: ("Frenzy", "Greatly raises attack power for a time."),
       5203: ("Gale", "Raises attack speed for a time."),
       5204: ("Bulwark", "Greatly raises defense for a time."),
       5205: ("Finisher", "Gathers all strength for a decisive blow.")}


def add_skills():
    path = os.path.join(XL, "SkillData.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    existing = {ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)}

    n = 0
    for sid, _kor, _desc, effect, stat, value, dur, cool, vfx in SKILLS:
        if sid in existing:
            continue
        row = {
            "DataId": sid,
            "Name": f"Skill{sid}NameText",
            "Description": f"Skill{sid}DescText",
            "IconLabel": f"Skill{sid}",
            "Category": "Combat",
            "SkillActiveType": "Repeat",
            "MaxLevel": 5,
            "MemoryCapacity": 2,
            "Source": "Unlimited",
            "TriggerType": "Active",
            "ProcChance": 100,
            "Cooldown": cool,
            "ActiveContext": "InBattle",
            "Effect": effect,
            "EffectStat": stat,
            "EffectValue": value,
            "EffectDuration": dur,
            "EffectVfxLabel": vfx,
        }
        ws.append([row.get(h) for h in head])
        n += 1
    wb.save(path)
    print(f"SkillData.xlsx — 액티브 {n}종 추가")


def add_localize():
    path = os.path.join(XL, "LocalizationData.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    have = {ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)}
    n = 0
    for sid, kor, desc, *_ in SKILLS:
        e_name, e_desc = ENG[sid]
        for key, k, e in ((f"Skill{sid}NameText", kor, e_name),
                          (f"Skill{sid}DescText", desc, e_desc)):
            if key in have:
                continue
            ws.append([key, k, e])
            n += 1
    wb.save(path)
    print(f"LocalizationData.xlsx — 키 {n}개 추가")


if __name__ == "__main__":
    add_skills()
    add_localize()
