"""첫 두 미니몬의 기존 패시브 보상을 보존하며 수동 능력을 추가한다.

기본 실행은 dry-run, --apply일 때만 원본 엑셀 두 파일을 수정한다.
이후 Unity Tools/ParseAllLocal로 JSON을 재생성해야 한다.
"""
import argparse
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
EXCEL = ROOT / 'Assets/@Resources/Data/Excel'

UPGRADES = {
    6001: {
        'SkillActiveType': 'Repeat', 'TriggerType': 'Active', 'RoleTypes': 'Def',
        'Effect': 'Heal', 'EffectStat': 'MaxHp', 'EffectValue': 18,
        'Cooldown': 18, 'EffectDuration': 0, 'EffectVfxLabel': 'HP Potion Ef',
    },
    6002: {
        'SkillActiveType': 'Repeat', 'TriggerType': 'Active', 'RoleTypes': 'Skill',
        'Effect': 'Buff', 'EffectStat': 'AtkSpeed', 'EffectValue': 30,
        'Cooldown': 18, 'EffectDuration': 8, 'EffectVfxLabel': 'Star_01_Anim_13_a',
    },
}
DESCRIPTIONS = {
    'Skill6001DescText': (
        '기본 효과: 단계당 최대 체력 +30.\n'
        '수동 사용: 최대 체력의 18% × 능력 단계만큼 회복한다.\n'
        '재사용 18초(미니몬 능력끼리 공유).\n'
        '동행 중 임시 사용 가능. 공명도 100%에서 영구 습득.',
        'Passive: +30 maximum HP per ability level.\n'
        'Manual use: Restore 18% of maximum HP per ability level.\n'
        'Cooldown: 18s, shared by minion abilities.\n'
        'Available temporarily while accompanying. Permanently learned at 100% resonance.'),
    'Skill6002DescText': (
        '기본 효과: 단계당 공격속도 +10%p.\n'
        '수동 사용: 8초 동안 공격속도가 30%p × 능력 단계만큼 오른다.\n'
        '재사용 18초(미니몬 능력끼리 공유).\n'
        '동행 중 임시 사용 가능. 공명도 100%에서 영구 습득.',
        'Passive: +10 percentage points of attack speed per ability level.\n'
        'Manual use: Gain 30 percentage points of attack speed per ability level for 8 seconds.\n'
        'Cooldown: 18s, shared by minion abilities.\n'
        'Available temporarily while accompanying. Permanently learned at 100% resonance.'),
}
OLD_LURE = '당신는(은) 매력발산!를(을) 사용했다!'
NEW_LURE = '매력을 발산해 미니몬의 관심을 끌었다!'


def index_sheet(ws, start):
    columns = {cell.value: cell.column for cell in ws[1] if cell.value is not None}
    rows = {}
    for row in range(start, ws.max_row + 1):
        key = ws.cell(row, 1).value
        if key is not None:
            assert key not in rows, f'중복 키: {key}'
            rows[key] = row
    return columns, rows


def upgrade(apply=False):
    books = {name: openpyxl.load_workbook(EXCEL / f'{name}.xlsx')
             for name in ('SkillData', 'LocalizationData')}
    changed = {name: [] for name in books}

    def patch(name, row, column, value):
        cell = books[name].active.cell(row, column)
        if cell.value != value:
            changed[name].append((cell.coordinate, cell.value, value))
            cell.value = value

    try:
        ws = books['SkillData'].active
        columns, rows = index_sheet(ws, 4)
        for sid, updates in UPGRADES.items():
            row = rows[sid]
            before = {key: ws.cell(row, column).value for key, column in columns.items()}
            stat, bonus = ('MaxHp', 30) if sid == 6001 else ('AtkSpeed', 10)
            expected = {'Category': 'Memory', 'StatType': stat, 'StatBonus': bonus,
                        'MaxLevel': 2, 'MemoryCapacity': 2, 'Source': 'Minion',
                        'RequiredMinionId': sid - 4000, 'Description': f'Skill{sid}DescText'}
            assert all(before[key] == value for key, value in expected.items()), (sid, before)
            for key, value in updates.items():
                patch('SkillData', row, columns[key], value)
            # 명시한 수동 능력 필드 외에는 이름/아이콘/출처/성장 등 모두 보존한다.
            assert all(ws.cell(row, column).value == before[key]
                       for key, column in columns.items() if key not in updates), sid

        ws = books['LocalizationData'].active
        columns, rows = index_sheet(ws, 2)
        for key, values in DESCRIPTIONS.items():
            for language, value in zip(('Kor', 'Eng'), values):
                patch('LocalizationData', rows[key], columns[language], value)
        assert ws.cell(rows['TryLureText'], columns['Kor']).value in (OLD_LURE, NEW_LURE)
        patch('LocalizationData', rows['TryLureText'], columns['Kor'], NEW_LURE)

        for name, changes in changed.items():
            print(f'{name}: {len(changes)}개 셀 변경, apply={apply}')
            for coordinate, before, after in changes:
                print(f'  {coordinate}: {before!r} -> {after!r}')
        if apply:
            for name, changes in changed.items():
                if changes:
                    books[name].save(EXCEL / f'{name}.xlsx')
    finally:
        for book in books.values():
            book.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    upgrade(parser.parse_args().apply)
