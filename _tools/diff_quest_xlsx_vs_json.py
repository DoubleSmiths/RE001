"""
Diff QuestData.xlsx vs the JSON files. Reports any drift.
"""
import json
from pathlib import Path
import openpyxl

ROOT = Path(r"C:\stopresent\Unity\RE001\Assets\@Resources\Data")
XLSX = ROOT / "Excel" / "QuestData.xlsx"
JSON_DIR = ROOT / "JsonData"

wb = openpyxl.load_workbook(XLSX, data_only=False)


def norm(v):
    if v is None or v == "":
        return None
    if isinstance(v, str):
        s = v.strip()
        if s == "":
            return None
        if s.upper() == "TRUE":
            return True
        if s.upper() == "FALSE":
            return False
        return s
    return v


def rows_dict(sheet_name, key_cols):
    ws = wb[sheet_name]
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    out = {}
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if row[0] is None:
            continue
        d = dict(zip(header, row))
        key = tuple(d[k] for k in key_cols)
        out[key] = d
    return out


def jload(name):
    p = JSON_DIR / f"{name}.json"
    return json.loads(p.read_text(encoding="utf-8-sig"))


# ── Quest ─────────────────────────────────────────
print("=== Quest ===")
xl_q = rows_dict("Quest", ["QuestID"])
js_q = {(q["QuestID"],): q for q in jload("QuestData")["quests"]}

# Map enum index ↔ name (as written in QuestManager / Define.cs)
COMPLETE_TYPE = ["ITEM_ACQUIRE", "ITEM_COUNT", "ANIM_END", "MULTI", "BTN_CLICK"]

for key in sorted(set(xl_q) | set(js_q)):
    qid = key[0]
    x = xl_q.get(key)
    j = js_q.get(key)
    if not x or not j:
        print(f"  Q{qid}: only in {'JSON' if j else 'Excel'}")
        continue
    fields = [
        ("CompleteType",   COMPLETE_TYPE[j["CompleteType"]],          x["CompleteType"]),
        ("CompleteTarget", j["CompleteTarget"] or None,               norm(x["CompleteTarget"])),
        ("CompleteCount",  j["CompleteCount"],                        norm(x["CompleteCount"]) or 0),
        ("FollowItemID",   j["FollowItemID"] or None,                 norm(x["FollowItemID"])),
        ("FollowItemName", j["FollowItemName"] or None,               norm(x["FollowItemName"])),
        ("FollowCount",    j["FollowCount"],                          norm(x["FollowCount"]) or 0),
        ("AnimationID",    j["AnimationID"] or None,                  norm(x["AnimationID"])),
        ("AnimBlocking",   j["AnimBlocking"],                         bool(norm(x["AnimBlocking"])) if norm(x["AnimBlocking"]) is not None else False),
        ("DialogueID",     j["DialogueID"] or None,                   norm(x["DialogueID"])),
    ]
    diffs = []
    for name, jval, xval in fields:
        # FollowItemID/CompleteTarget JSON은 string, Excel은 int일 수 있음 → str로 비교
        if name in ("CompleteTarget", "FollowItemID") and jval is not None and xval is not None:
            jval, xval = str(jval), str(xval)
        if jval != xval:
            diffs.append(f"{name}: JSON={jval!r}  XLSX={xval!r}")
    if diffs:
        print(f"  Q{qid} DRIFT:")
        for d in diffs:
            print(f"    - {d}")
    else:
        print(f"  Q{qid} OK")


# ── Quest_Condition ───────────────────────────────
print("\n=== Quest_Condition ===")
COND_TYPE = ["BTN_CLICK", "ITEM_ACQUIRE", "ITEM_COUNT"]
xl_c = rows_dict("Quest_Condition", ["QuestID", "Order"])
js_c = {(c["QuestID"], c["Order"]): c for c in jload("QuestConditionData")["conditions"]}
for k in sorted(set(xl_c) | set(js_c)):
    x = xl_c.get(k)
    j = js_c.get(k)
    if not x or not j:
        print(f"  {k}: only in {'JSON' if j else 'Excel'}: {(j or x)}")
        continue
    j_type = COND_TYPE[j["ConditionType"]]
    j_target = str(j["Target"])
    x_type = norm(x["CompleteType"])
    x_target = str(norm(x["CompleteTarget"]))
    if j_type != x_type or j_target != x_target:
        print(f"  {k} DRIFT: JSON={(j_type, j_target)}  XLSX={(x_type, x_target)}")


# ── Quest_Notify ──────────────────────────────────
print("\n=== Quest_Notify ===")
xl_n = rows_dict("Quest_Notify", ["QuestID", "Order"])
js_n = {(n["QuestID"], n["Order"]): n for n in jload("QuestNotifyData")["notifies"]}
for k in sorted(set(xl_n) | set(js_n)):
    x = xl_n.get(k)
    j = js_n.get(k)
    if not x or not j:
        print(f"  {k}: only in {'JSON' if j else 'Excel'}: {(j or x)}")
        continue
    if j["BtnID"] != norm(x["BtnID"]):
        print(f"  {k} DRIFT: JSON={j['BtnID']!r}  XLSX={norm(x['BtnID'])!r}")


# ── Quest_Unlock ──────────────────────────────────
print("\n=== Quest_Unlock ===")
UNLOCK_TYPE = ["UI_BUTTON", "SUBQUEST_START", "ITEM_REWARD", "ITEM_REMOVE"]
xl_u = rows_dict("Quest_Unlock", ["QuestID", "TargetID"])
# 키를 (QuestID, str(TargetID))로 정규화 (Excel int vs JSON str 혼용 흡수)
xl_u = {(qid, str(tid)): v for (qid, tid), v in xl_u.items()}
js_u = {(u["QuestID"], str(u["TargetID"])): u for u in jload("QuestUnlockData")["unlocks"]}
for k in sorted(set(xl_u) | set(js_u)):
    x = xl_u.get(k)
    j = js_u.get(k)
    if not x or not j:
        print(f"  {k}: only in {'JSON' if j else 'Excel'}: {(j or x)}")
        continue
    j_type = UNLOCK_TYPE[j["UnlockType"]]
    x_type = norm(x["UnlockType"])
    if j_type != x_type:
        print(f"  {k} DRIFT: JSON UnlockType={j_type}  XLSX={x_type}")
