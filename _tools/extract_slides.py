import sys
import io
from pptx import Presentation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

path = sys.argv[1] if len(sys.argv) > 1 else "Documents/기획안/기획서v2_04.16.pptx"
prs = Presentation(path)

for i, slide in enumerate(prs.slides, 1):
    print(f"\n===== SLIDE {i} =====")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text:
                    print(text)
        if shape.shape_type == 13:
            try:
                print(f"[IMG: {shape.name}]")
            except Exception:
                pass
