"""직업 능력 선택과 한영 설명을 보강한다. 기본 dry-run, --apply로만 엑셀 저장.

기존/최종 값만 허용하며 SkillData 5004의 세 셀 외 수치는 보존한다.
동적 설명 신규 키는 class_ability_text.TRANSLATIONS가 소유한다.
적용 후 Unity Tools/ParseAllLocal로 JSON을 재생성해야 한다.
"""
import argparse
from copy import copy
from pathlib import Path

import openpyxl

EXCEL = Path(__file__).resolve().parents[1] / 'Assets/@Resources/Data/Excel'

# (한글 이름, 기존 영문 이름, 최종 영문 이름, 기존 한글 설명, 기존 영문 설명, 최종 한글 설명, 최종 영문 설명)
CLASS_TEXT = {
    5004: ('강철의 기억', 'Memory of Steel', 'Memory of Steel',
           '최대 체력을 증가시킨다.', 'Increases max HP.',
           '최대 체력과 방어력이 오른다.', 'Increases max HP and defense.'),
    5007: ('전투의 투지', 'Battle Spirit', 'Battle Spirit',
           '전투 시작 시 일정 시간 물리 공격력이 증가한다.', 'Physical attack rises for a while at battle start.',
           '전투 시작 시 일정 시간 물리 공격력이 증가한다.',
           'Increases physical attack briefly when entering a battle area.'),
    5100: ('선제의 감각', '선제의 감각', 'First Strike Instinct',
           '전투 시작 시 60% 확률로 공격력이 잠시 오른다', '전투 시작 시 60% 확률로 공격력이 잠시 오른다',
           '전투 시작 시 60% 확률로 공격력이 잠시 오른다',
           'A 60% chance to briefly increase physical attack when entering a battle area.'),
    5105: ('연쇄 각성', '연쇄 각성', 'Chain Awakening',
           '연속 3회 치명타 시 큰 피해를 추가한다', '연속 3회 치명타 시 큰 피해를 추가한다',
           '연속 3회 치명타 시 큰 피해를 추가한다',
           'Deals bonus damage after three consecutive critical hits.'),
    5109: ('그림자 걸음', '그림자 걸음', 'Shadow Step',
           '연속 3회 회피 시 회피율이 크게 오른다', '연속 3회 회피 시 회피율이 크게 오른다',
           '연속 3회 회피 시 회피율이 크게 오른다',
           'Briefly increases evasion after three consecutive dodges.'),
    5113: ('불굴', '불굴', 'Indomitable',
           '치명상을 입을 때 40% 확률로 체력을 회복하며 버틴다', '치명상을 입을 때 40% 확률로 체력을 회복하며 버틴다',
           '치명 피해를 받기 직전 40% 확률로 체력을 회복한다.',
           'A 40% chance to restore HP just before taking an otherwise fatal hit.'),
    5116: ('각성', '각성', 'Awakening',
           '18초마다 공격력이 잠시 오른다', '18초마다 공격력이 잠시 오른다',
           '18초마다 공격력이 잠시 오른다',
           'Briefly increases physical attack every 18 seconds while in battle.'),
    5121: ('잔상', '잔상', 'Afterimage',
           '17초마다 회피율이 잠시 오른다', '17초마다 회피율이 잠시 오른다',
           '17초마다 회피율이 잠시 오른다',
           'Briefly increases evasion every 17 seconds while in battle.'),
    5125: ('강인함', '강인함', 'Fortitude',
           '최대 체력이 오른다', '최대 체력이 오른다',
           '최대 체력이 오른다', 'Increases max HP.'),
    5131: ('민첩', '민첩', 'Agility',
           '회피율이 오른다', '회피율이 오른다',
           '회피율이 오른다', 'Increases evasion.'),
    5139: ('각인', '각인', 'Imprint',
           '확정 포획에 필요한 횟수가 줄어든다', '확정 포획에 필요한 횟수가 줄어든다',
           '확정 포획에 필요한 횟수가 줄어든다',
           'Reduces the number of attempts needed for a guaranteed capture.'),
    5142: ('연계 감각', '연계 감각', 'Combo Instinct',
           '콤보 게이지가 더 오래 유지된다', '콤보 게이지가 더 오래 유지된다',
           '콤보 유지 시간이 늘어난다.', 'Extends the time before a combo expires.'),
}


def index_sheet(ws, start):
    columns = {cell.value: cell.column for cell in ws[1]}
    rows = {}
    for row in ws.iter_rows(min_row=start):
        key = row[0].value
        if key is None:
            continue
        assert key not in rows, f'{ws.title} 중복 키: {key}'
        rows[key] = row[0].row
    return columns, rows


def prepare(books):
    """메모리에서만 변경 계획을 적용한다. 같은 책에 두 번 호출하면 두 번째는 변경 0."""
    from class_ability_text import TRANSLATIONS

    changes = {name: {} for name in books}
    before = {name: {(ws.title, cell.coordinate): (cell.value, copy(cell._style))
                     for ws in book for row in ws for cell in row} for name, book in books.items()}
    original_rows = {name: book.active.max_row for name, book in books.items()}
    added_keys = []

    def patch(name, row, column, old, final):
        cell = books[name].active.cell(row, column)
        assert cell.value == old or cell.value == final, f'{name} {cell.coordinate} 예상 밖 값: {cell.value!r}'
        if cell.value != final:
            changes[name][(books[name].active.title, cell.coordinate)] = (cell.value, final)
            cell.value = final

    skill_ws = books['SkillData'].active
    columns, skills = index_sheet(skill_ws, 4)
    # Unity 파서는 0-based 28/29열을 쉼표로 나누므로 헤더 위치까지 검증한다.
    assert columns['StatTypes'] == 29 and columns['StatBonuses'] == 30
    class_ids = {sid for sid, row in skills.items() if skill_ws.cell(row, columns['Source']).value == 'Class'}
    assert class_ids == set(CLASS_TEXT), f'Class 능력 목록 변경: {class_ids}'
    for sid, stat, bonus, memory in ((5004, 'MaxHp', 50, None), (5125, 'MaxHp', 120, 2)):
        row = skills[sid]
        assert skill_ws.cell(row, columns['StatType']).value == stat
        assert skill_ws.cell(row, columns['StatBonus']).value == bonus
        if memory is not None:
            assert skill_ws.cell(row, columns['MemoryCapacity']).value == memory
    for column, old, final in (('MemoryCapacity', 4, 2), ('StatTypes', None, 'MaxHp,Def'), ('StatBonuses', None, '50,5')):
        patch('SkillData', skills[5004], columns[column], old, final)

    loc_ws = books['LocalizationData'].active
    loc_columns, loc_rows = index_sheet(loc_ws, 2)
    for sid, (kor_name, old_eng_name, eng_name, old_kor_desc, old_eng_desc, kor_desc, eng_desc) in CLASS_TEXT.items():
        name_key, desc_key = f'Skill{sid}NameText', f'Skill{sid}DescText'
        assert skill_ws.cell(skills[sid], columns['Name']).value == name_key
        assert skill_ws.cell(skills[sid], columns['Description']).value == desc_key
        assert loc_ws.cell(loc_rows[name_key], loc_columns['Kor']).value == kor_name
        patch('LocalizationData', loc_rows[name_key], loc_columns['Eng'], old_eng_name, eng_name)
        patch('LocalizationData', loc_rows[desc_key], loc_columns['Kor'], old_kor_desc, kor_desc)
        patch('LocalizationData', loc_rows[desc_key], loc_columns['Eng'], old_eng_desc, eng_desc)

    assert not set(TRANSLATIONS) & {f'Skill{sid}{part}Text' for sid in CLASS_TEXT for part in ('Name', 'Desc')}
    for key, pair in TRANSLATIONS.items():
        assert isinstance(pair, (tuple, list)) and len(pair) == 2 and all(isinstance(v, str) and v for v in pair), key
        if key in loc_rows:
            row = loc_rows[key]
            assert tuple(loc_ws.cell(row, loc_columns[language]).value for language in ('Kor', 'Eng')) == tuple(pair), key
            continue
        row = loc_ws.max_row + 1
        for column, value in enumerate((key, *pair), 1):
            cell = loc_ws.cell(row, column, value)
            cell._style = copy(loc_ws.cell(original_rows['LocalizationData'], column)._style)
            changes['LocalizationData'][(loc_ws.title, cell.coordinate)] = (None, value)
        loc_rows[key] = row
        added_keys.append(key)

    assert skill_ws.max_row == original_rows['SkillData']
    assert loc_ws.max_row == original_rows['LocalizationData'] + len(added_keys)
    for name, cells in before.items():
        for (sheet, coordinate), (value, style) in cells.items():
            expected = changes[name].get((sheet, coordinate), (value, value))[1]
            cell = books[name][sheet][coordinate]
            assert cell.value == expected and cell._style == style, f'허용 외 변경: {name} {coordinate}'
    return changes, added_keys


def upgrade(apply=False):
    books = {name: openpyxl.load_workbook(EXCEL / f'{name}.xlsx') for name in ('SkillData', 'LocalizationData')}
    try:
        changes, added_keys = prepare(books)
        for name, cells in changes.items():
            print(f'{name}: {len(cells)}개 셀 변경, apply={apply}')
        print(f'동적 설명 키: {len(added_keys)}개 추가')
        # 파일 쓰기 없이 같은 워크북을 다시 처리해 멱등성을 검증한다.
        repeated, repeated_keys = prepare(books)
        assert not repeated_keys and not any(repeated.values()), '동일 데이터 재실행 변경 발생'
        if apply:
            for name, cells in changes.items():
                if cells:
                    books[name].save(EXCEL / f'{name}.xlsx')
    finally:
        for book in books.values():
            book.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply', action='store_true')
    upgrade(parser.parse_args().apply)
