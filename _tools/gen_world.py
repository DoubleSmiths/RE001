# -*- coding: utf-8 -*-
"""세계 구조 재생성 — 마을 5 × 지역 5 × 던전 5층 + 몬스터별 고유 재료.

해결하는 문제
  ③ 마을이 아스달 하나뿐
  ④ 중분류(지역)가 3개, 소분류(던전)가 4개뿐 — 요구는 마을당 지역 5+, 지역당 던전 5+
  ⑩ 몬스터 503종이 드롭 재료를 2종만 돌려씀 → 몬스터 외형과 재료가 이어지지 않음

설계
  마을 5개가 60레벨씩 담당(1~60 / 61~120 / … / 241~300).
  마을당 지역 5개(각 12레벨), 지역당 던전 5층(각 2~3레벨) = **125개 전투 노드**.
  방치형이라 해금이 잦은 편이 낫다 — 2~3레벨마다 새 층이 열린다.

  재료는 **몬스터 원형(아키타입) 30종에 1:1로** 만든다.
  503종 각각에 고유 재료를 만들면 인벤토리가 무의미해지고, 2종만 쓰면 외형과 이어지지 않는다.
  원형이 곧 화면에 보이는 몬스터 이미지이므로 이 단위가 정확히 맞다.
  (슬라임 → 슬라임 젤리, 숲거미 → 거미줄 …)

ID 체계
  마을     11000 + t*1000                (t=0..4)
  지역     21000 + t*1000 + r*100        (r=0..4)
  던전     지역ID + 10*(d+1)             (d=0..4)
  재료     11001 + 아키타입 index

실행: python _tools/gen_world.py → Unity에서 Tools/ParseAllLocal
주의: 여정 시트는 **전면 재작성**한다(기존 3지역 30층 구조를 대체).
"""
import os
import openpyxl

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
EXCEL = os.path.join(BASE, "Assets", "@Resources", "Data", "Excel")

TOWNS = [
    dict(key="Asdal",     kor="아스달 마을",     eng="Asdal Village",     lv0=1,   sprite="Background_Village1"),
    dict(key="Windport",  kor="바람서린 항구",   eng="Windrime Port",     lv0=61,  sprite="Background_Village1"),
    dict(key="Ashmine",   kor="잿빛 광산촌",     eng="Ashen Mining Town", lv0=121, sprite="Background_Village1"),
    dict(key="Frostkeep", kor="서리안개 요새",   eng="Frostmist Keep",    lv0=181, sprite="Background_Village1"),
    dict(key="Dawnspire", kor="여명의 성도",     eng="Dawnspire",         lv0=241, sprite="Background_Village1"),
]

# 마을별 지역 5개. (키, 한글, 영문) — 원형 스프라이트 계열은 REGION_ARCH로 매핑
REGIONS = [
    [("BreezingWindForest", "신들바람 숲",   "Breezing Wind Forest"),
     ("WhisperingCave",     "속삭이는 동굴", "Whispering Cave"),
     ("RuinedSanctum",      "폐허가 된 성소", "Ruined Sanctum"),
     ("SunkenGrove",        "가라앉은 숲",   "Sunken Grove"),
     ("HollowBurrow",       "텅 빈 굴",      "Hollow Burrow")],
    [("SaltMarsh",          "소금 늪지",     "Salt Marsh"),
     ("TidalCavern",        "밀물 동굴",     "Tidal Cavern"),
     ("BrokenPier",         "부서진 잔교",   "Broken Pier"),
     ("StormCliff",         "폭풍 절벽",     "Storm Cliff"),
     ("DriftwoodWood",      "표류목 숲",     "Driftwood Wood")],
    [("CinderPit",          "잿불 구덩이",   "Cinder Pit"),
     ("DeepShaft",          "깊은 수직굴",   "Deep Shaft"),
     ("RustHall",           "녹슨 회랑",     "Rust Hall"),
     ("EmberVein",          "불씨 광맥",     "Ember Vein"),
     ("CollapsedTunnel",    "무너진 갱도",   "Collapsed Tunnel")],
    [("FrostWood",          "서리 수림",     "Frost Wood"),
     ("GlacierCave",        "빙하 동굴",     "Glacier Cave"),
     ("SilentChapel",       "고요한 예배당", "Silent Chapel"),
     ("WhiteWaste",         "새하얀 황야",   "White Waste"),
     ("FrozenVault",        "얼어붙은 금고", "Frozen Vault")],
    [("DawnGarden",         "여명의 정원",   "Dawn Garden"),
     ("VoidStair",          "공허의 계단",   "Void Stair"),
     ("StarSanctum",        "별의 성소",     "Star Sanctum"),
     ("LastRampart",        "마지막 성벽",   "Last Rampart"),
     ("ThroneOfDawn",       "여명의 왕좌",   "Throne of Dawn")],
]

DEPTH_KOR = ["입구", "2층", "3층", "4층", "심층"]
DEPTH_ENG = ["Entrance", "F2", "F3", "F4", "Depths"]

# 지역 → 사용할 몬스터 원형 계열(기존 생성 스프라이트 3계열을 순환)
ARCH_FAMILY = ["BreezingWindForest", "WhisperingCave", "RuinedSanctum"]

# 원형 30종 ↔ 재료 (몬스터 외형과 이어지도록 이름을 맞춘다)
MATERIALS = [
    # 숲 계열 0~9
    ("슬라임 젤리", "Slime Jelly"), ("거미 실크", "Spider Silk"), ("늑대 송곳니", "Wolf Fang"),
    ("독버섯 갓", "Toadstool Cap"), ("바람 정수", "Wind Essence"), ("멧돼지 가죽", "Boar Hide"),
    ("까마귀 깃털", "Crow Feather"), ("질긴 덩굴", "Tough Vine"), ("여우 꼬리털", "Fox Fur"),
    ("말벌 침", "Hornet Sting"),
    # 동굴 계열 10~19
    ("박쥐 날개막", "Bat Membrane"), ("두더지 발톱", "Mole Claw"), ("광부의 유품", "Miner's Relic"),
    ("수정 파편", "Crystal Shard"), ("지네 외피", "Centipede Shell"), ("젖은 이끼", "Damp Moss"),
    ("석상 조각", "Statue Fragment"), ("그림자 응결물", "Shadow Clot"), ("도롱뇽 비늘", "Salamander Scale"),
    ("발광 점액", "Glow Slime"),
    # 성소 계열 20~29
    ("망령의 재", "Revenant Ash"), ("기사의 문장", "Knight's Crest"), ("타락한 성물", "Tainted Relic"),
    ("키메라 갈기", "Chimera Mane"), ("가고일 돌조각", "Gargoyle Stone"), ("수호자의 핵", "Warden Core"),
    ("환영의 조각", "Phantom Sliver"), ("오염된 정수", "Tainted Essence"), ("제단의 파편", "Altar Fragment"),
    ("공허 결정", "Void Crystal"),
]
MAT_ID0 = 11001

GRADE_BY_TIER = ["Basic", "Basic", "Abnormal", "Abnormal", "Rare",
                 "Rare", "Exceptional", "Exceptional", "Legendary", "Mystic"]


def town_of(level):
    return min(4, max(0, (level - 1) // 60))


def build_nodes():
    """(마을들, 지역들, 던전들) — 각 원소는 dict."""
    towns, regions, dungeons = [], [], []
    for t, town in enumerate(TOWNS):
        tid = 11000 + t * 1000
        towns.append(dict(id=tid, key=town["key"], kor=town["kor"], eng=town["eng"],
                          sprite=town["sprite"], lv=town["lv0"]))
        for r, (rkey, rkor, reng) in enumerate(REGIONS[t]):
            rid = 21000 + t * 1000 + r * 100
            rlv0 = town["lv0"] + r * 12
            regions.append(dict(id=rid, town=tid, key=rkey, kor=rkor, eng=reng, lv=rlv0,
                                arch=ARCH_FAMILY[(t * 5 + r) % len(ARCH_FAMILY)]))
            for d in range(5):
                did = rid + 10 * (d + 1)
                # 12레벨을 5층에 나눔 → 0,2,4,7,9 오프셋
                dlv = min(300, rlv0 + [0, 2, 4, 7, 9][d])
                dungeons.append(dict(id=did, region=rid, key=f"{rkey}D{d+1}",
                                     kor=f"{rkor} {DEPTH_KOR[d]}", eng=f"{reng} {DEPTH_ENG[d]}",
                                     lv=dlv, arch=ARCH_FAMILY[(t * 5 + r) % len(ARCH_FAMILY)],
                                     tier=min(9, (dlv - 1) // 30)))
    return towns, regions, dungeons


def write_travel(towns, regions, dungeons):
    p = os.path.join(EXCEL, "TravelData.xlsx")
    wb = openpyxl.load_workbook(p)
    node, cond = wb["TravelNode"], wb["TravelCondition"]

    for ws in (node, cond):
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)      # 헤더만 남기고 전면 재작성

    r = 2
    for t in towns:
        node.cell(r, 1, t["id"]); node.cell(r, 2, 0); node.cell(r, 3, "TOWN")
        node.cell(r, 4, f'Place{t["key"]}Text'); node.cell(r, 5, t["sprite"])
        node.cell(r, 6, "Village1")
        node.cell(r, 7, f'BTN_MAP_{t["key"].upper()}'); node.cell(r, 8, f'BTN_ENTER_{t["key"].upper()}')
        r += 1
    for g in regions:
        node.cell(r, 1, g["id"]); node.cell(r, 2, 0); node.cell(r, 3, "EXPLORE")
        node.cell(r, 4, f'{g["key"]}NameText'); node.cell(r, 5, f'{g["arch"]}_Img')
        node.cell(r, 7, f'BTN_MAP_{g["key"].upper()}')
        r += 1
    for d in dungeons:
        node.cell(r, 1, d["id"]); node.cell(r, 2, d["region"]); node.cell(r, 3, "EXPLORE")
        node.cell(r, 4, f'{d["key"]}Text'); node.cell(r, 5, f'{d["arch"]}_Entry_Img')
        node.cell(r, 6, "Battle")
        node.cell(r, 7, f'BTN_MAP_{d["key"].upper()}'); node.cell(r, 8, f'BTN_STAGE_{d["key"].upper()}')
        r += 1

    c = 2
    for g in regions:
        if g["lv"] <= 1: continue
        cond.cell(c, 1, g["id"]); cond.cell(c, 2, "LEVEL_REQUIRED")
        cond.cell(c, 3, str(g["lv"])); cond.cell(c, 4, g["lv"])
        cond.cell(c, 5, f'레벨 {g["lv"]} 이상이어야 갈 수 있습니다.')
        c += 1
    for d in dungeons:
        if d["lv"] <= 1: continue
        cond.cell(c, 1, d["id"]); cond.cell(c, 2, "LEVEL_REQUIRED")
        cond.cell(c, 3, str(d["lv"])); cond.cell(c, 4, d["lv"])
        cond.cell(c, 5, f'레벨 {d["lv"]} 이상이어야 들어갈 수 있습니다.')
        c += 1
    wb.save(p)
    print(f"TravelData: 마을 {len(towns)} / 지역 {len(regions)} / 던전 {len(dungeons)}")


def write_materials():
    """InventoryItem 시트에 아키타입 재료 30종 추가(헤더 1줄, 데이터 2행부터)."""
    p = os.path.join(EXCEL, "ItemData.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb["InventoryItem"]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    have = {ws.cell(r, h["DataId"]).value for r in range(2, ws.max_row + 1)}
    row = ws.max_row + 1
    added = 0
    for i, (kor, eng) in enumerate(MATERIALS):
        did = MAT_ID0 + i
        if did in have: continue
        fam = ARCH_FAMILY[i // 10]
        ws.cell(row, h["DataId"], did)
        ws.cell(row, h["Name"], f"Mat{did}NameText")
        ws.cell(row, h["Grade"], "Basic" if i % 10 < 6 else "Abnormal")
        ws.cell(row, h["Description"], f"Mat{did}DescText")
        ws.cell(row, h["SpriteName"], f"Mat_{fam}_{i % 10}")
        ws.cell(row, h["FoundLocation"], f"{fam}NameText")
        row += 1; added += 1
    wb.save(p)
    print(f"ItemData: 재료 +{added}종")


def rewire_creatures(dungeons):
    """503종을 125개 던전에 재배치 + 원형에 맞는 재료를 드롭하도록 연결."""
    p = os.path.join(EXCEL, "CreatureData.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb[wb.sheetnames[0]]
    h = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}

    rows = [r for r in range(4, ws.max_row + 1) if ws.cell(r, h["DataId"]).value is not None]
    per = max(1, len(rows) // len(dungeons))
    for i, r in enumerate(rows):
        d = dungeons[min(len(dungeons) - 1, i // per)]
        ws.cell(r, h["Habitat"], f'{d["key"]}Text')
        # 원형: 애니메이터가 이미 Mon_{id}이므로 IconLabel(=원형 스프라이트)만 지역 계열에 맞춘다
        did = int(ws.cell(r, h["DataId"]).value)
        arch_idx = abs(did) % 10
        ws.cell(r, h["IconLabel"], f'Mon_{d["arch"]}_{arch_idx}')
        # 재료: 원형과 1:1 (계열 index × 10 + 원형 index)
        fam_idx = ARCH_FAMILY.index(d["arch"])
        ws.cell(r, h["DropItemId"], MAT_ID0 + fam_idx * 10 + arch_idx)
    wb.save(p)
    print(f"CreatureData: {len(rows)}종을 던전 {len(dungeons)}개에 재배치 (노드당 약 {per}종), 재료 연결")


def write_localization(towns, regions, dungeons):
    p = os.path.join(EXCEL, "LocalizationData.xlsx")
    wb = openpyxl.load_workbook(p)
    ws = wb[wb.sheetnames[0]]
    have = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
    row = ws.max_row + 1
    n = 0

    def add(key, kor, eng):
        nonlocal row, n
        if key in have: return
        ws.cell(row, 1, key); ws.cell(row, 2, kor); ws.cell(row, 3, eng)
        have.add(key); row += 1; n += 1

    for t in towns:
        add(f'Place{t["key"]}Text', t["kor"], t["eng"])
    for g in regions:
        add(f'{g["key"]}NameText', g["kor"], g["eng"])
    for d in dungeons:
        add(f'{d["key"]}Text', d["kor"], d["eng"])
    for i, (kor, eng) in enumerate(MATERIALS):
        did = MAT_ID0 + i
        add(f"Mat{did}NameText", kor, eng)
        add(f"Mat{did}DescText", f"{kor}. 몬스터에게서 얻는 재료다.", f"{eng}. A crafting material.")
    wb.save(p)
    print(f"LocalizationData: +{n}개")


def main():
    towns, regions, dungeons = build_nodes()
    write_travel(towns, regions, dungeons)
    write_materials()
    rewire_creatures(dungeons)
    write_localization(towns, regions, dungeons)
    print("\n레벨 분포 샘플")
    for i in (0, 24, 49, 74, 99, 124):
        d = dungeons[i]
        print(f"  {d['kor']:24s} Lv{d['lv']:3d} (마을 {town_of(d['lv'])+1})")


if __name__ == "__main__":
    main()
