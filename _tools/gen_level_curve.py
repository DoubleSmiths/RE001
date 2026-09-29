# -*- coding: utf-8 -*-
"""LevelData.xlsx 생성 — 플레이어 1~300레벨 요구 경험치 곡선.

설계 근거(난이도/LOD):
  - Lv1~40 구간은 기존 밸런스(등비 ×1.2, Lv1=100 → Lv40=122,481)를 **그대로 보존**한다.
    이미 플레이 검증된 초반 체감을 건드리지 않기 위함.
  - Lv41부터는 등비 배율을 밴드별로 낮춘다. ×1.2를 300레벨까지 밀면 요구경험치가
    10^25에 도달해 어떤 몬스터 경험치로도 도달 불가능해진다(재미 규칙 위반).
  - 목표: 레벨당 처치 수를 초반 ~10킬에서 후반 ~60킬 사이로 유지.
    몬스터 TotalExp는 이 곡선에 맞춰 구역(티어)별로 등비 배치한다.

v2 재설계(ADR0030, 설계서 4.1): 위 곡선을 '원래 값'(BaseRequiredExp)으로 두고, 실제 RequiredExp = 원래 값 × m(L).
  m(L) 은 초반 거의 1배 → 중후반 최대 18배. 오프라인 정산만으로 이틀이면 만렙이 되던 수명을 2~5주로 늘린다.
  배수 곡선은 sim_economy.curve_mult 한 곳에만 있다. 원래 곡선으로 되돌리려면 그 배수를 1로.

실행: python _tools/gen_level_curve.py           # xlsx 저장
      python _tools/gen_level_curve.py --check   # 저장된 xlsx 가 이 공식과 같은지 검사(저장 안 함)
출력: Assets/@Resources/Data/Excel/LevelData.xlsx
      (JSON 반영은 Unity에서 Tools/ParseAllLocal. DataTransformer 는 A·B 열만 읽는다)
"""
import os
import sys

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from sim_economy import curve_mult  # noqa: E402

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "..", "Assets", "@Resources", "Data", "Excel", "LevelData.xlsx")

# (구간 끝 레벨, 배율) — 앞 구간부터 순서대로 적용
BANDS = [
    (40,  1.20),   # 튜토리얼~초반: 기존 곡선 보존
    (70,  1.12),   # 중반 진입
    (110, 1.09),
    (160, 1.07),
    (210, 1.055),
    (260, 1.045),
    (300, 1.035),  # 엔드게임: 완만하게
]


def rate_for(level):
    for end, r in BANDS:
        if level <= end:
            return r
    return BANDS[-1][1]


def build():
    rows = []
    exp = 100.0
    for lv in range(1, 301):
        if lv > 1:
            exp *= rate_for(lv)
        rows.append((lv, round(exp)))
    return rows


def build_v2():
    """(레벨, 실제 요구경험치, 원래 요구경험치)"""
    return [(lv, round(base * curve_mult(lv)), base) for lv, base in build()]


def main():
    rows = build_v2()

    if "--check" in sys.argv:
        ws = load_workbook(os.path.normpath(OUT), read_only=True).worksheets[0]
        saved = [tuple(r[:3]) for r in list(ws.iter_rows(values_only=True))[3:] if r[0] is not None]
        bad = [(a, b) for a, b in zip(saved, rows) if a != b]
        ok = not bad and len(saved) == len(rows)
        print("CHECK PASS" if ok else f"CHECK FAIL {len(bad)}행: {bad[:3]}")
        sys.exit(0 if ok else 1)

    wb = Workbook()
    ws = wb.active
    ws.title = "LevelData"
    ws.append(["Level", "RequiredExp", "BaseRequiredExp"])                      # Row0: 필드명
    ws.append(["int", "float", "float"])                                        # Row1: 타입
    ws.append(["레벨", "다음 레벨까지 필요 경험치", "원래 필요 경험치(v2 배수 전)"])  # Row2: 한글 설명
    for row in rows:                                                            # Row3+: 데이터
        ws.append(list(row))

    wb.save(os.path.normpath(OUT))
    print("saved:", os.path.normpath(OUT))

    # 곡선 확인용 요약
    for lv, exp, base in rows:
        if lv in (1, 10, 20, 40, 41, 70, 100, 150, 200, 250, 300):
            print(f"  Lv{lv:3d}  요구경험치 {exp:,}  (원래 {base:,} × {curve_mult(lv):.2f})")


if __name__ == "__main__":
    main()
